#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""V6.1 키워드별 본문/HTML 페이지 생성기.

사용:
    python make_sites_v6_1.py
    python make_sites_v6_1.py --input keywords.xlsx --output deploys --limit 100
"""
from __future__ import annotations

import argparse
import hashlib
import html
import random
import re
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl이 필요합니다: pip install openpyxl") from exc

BRAND = "체인지클린"
PHONE = "1688-1791"

INTRO = [
    "청소는 같은 평수라도 오염 상태와 구조에 따라 필요한 작업이 달라집니다.",
    "눈에 잘 띄는 바닥뿐 아니라 수납장 안쪽과 창틀처럼 놓치기 쉬운 구역까지 확인해야 합니다.",
    "현장 상태를 먼저 살핀 뒤 구역별 오염에 맞춰 작업 순서를 정하는 것이 중요합니다.",
    "입주나 이사를 앞둔 공간은 공사 분진과 생활 오염이 함께 남아 있는 경우가 많습니다.",
]
PROCESS = [
    ["현장 구조와 오염도 확인", "먼지 위쪽부터 아래쪽 순서로 제거", "주방·욕실 오염에 맞는 세정", "바닥 마감 및 누락 구역 점검"],
    ["고객 요청사항과 제외 범위 확인", "수납장·몰딩·창틀 건식 먼지 제거", "물때·기름때 집중 세척", "환기 후 최종 검수"],
    ["공간별 작업 동선 설정", "먼지와 폐기물 1차 정리", "오염 종류별 습식 작업", "마감 닦기와 사진 확인"],
]
CHECKS = [
    "엘리베이터와 주차 사용 가능 여부", "붙박이 가전 내부 청소 포함 여부", "스티커·보호필름 제거 범위",
    "곰팡이·니코틴·폐기물 등 특수오염 유무", "수도와 전기 사용 가능 여부", "입주 또는 이사 예정 시간",
]


def pick_rng(keyword: str) -> random.Random:
    seed = int(hashlib.sha256(keyword.encode("utf-8")).hexdigest()[:16], 16)
    return random.Random(seed)


def load_rows(path: Path) -> list[dict[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(v or "").strip() for v in next(ws.iter_rows(values_only=True))]
    rows = []
    for values in ws.iter_rows(values_only=True):
        row = {headers[i]: str(values[i] or "").strip() for i in range(min(len(headers), len(values)))}
        if row.get("keyword") and row.get("slug"):
            rows.append(row)
    return rows


def render_body(row: dict[str, str]) -> str:
    kw = row["keyword"]
    region = row.get("region", "해당 지역")
    neighborhood = row.get("neighborhood", "현장 인근")
    service = row.get("service", "청소")
    detail = row.get("detail", "구역별 오염 제거")
    prop = row.get("property_type", "공간")
    rng = pick_rng(kw)
    intro = rng.choice(INTRO)
    steps = rng.choice(PROCESS)
    checks = rng.sample(CHECKS, 4)
    focus = [
        f"{detail} 작업은 표면만 닦기보다 오염이 생긴 원인과 재질을 함께 확인합니다.",
        f"{prop} 구조에 따라 가구 안쪽, 걸레받이, 문틀, 배수구처럼 먼지가 모이는 지점이 달라집니다.",
        f"{region} {neighborhood} 현장은 주차와 작업 동선을 미리 확인하면 당일 진행이 한결 원활합니다.",
    ]
    rng.shuffle(focus)
    li_steps = "".join(f"<li><strong>{i+1}단계</strong> — {html.escape(s)}</li>" for i, s in enumerate(steps))
    li_checks = "".join(f"<li>{html.escape(s)}</li>" for s in checks)
    paragraphs = "".join(f"<p>{html.escape(p)}</p>" for p in focus)
    return f"""
<article class="content">
  <header class="hero">
    <p class="eyebrow">{html.escape(region)} {html.escape(service)} 현장 안내</p>
    <h1>{html.escape(kw)}</h1>
    <p>{html.escape(intro)}</p>
  </header>
  <section>
    <h2>{html.escape(neighborhood)} {html.escape(prop)} 청소에서 확인할 부분</h2>
    {paragraphs}
  </section>
  <section>
    <h2>작업 진행 순서</h2>
    <ol>{li_steps}</ol>
  </section>
  <section>
    <h2>견적 전 확인사항</h2>
    <ul>{li_checks}</ul>
    <p>평수만으로 확정하기 어려운 특수오염이나 폐기물이 있다면 사진을 함께 전달하는 것이 정확합니다.</p>
  </section>
  <section class="cta">
    <h2>{html.escape(BRAND)} 상담</h2>
    <p>현장 위치, 공간 유형, 평수, 희망 날짜, 특이 오염을 알려주시면 작업 범위를 확인합니다.</p>
    <p><a href="tel:{PHONE.replace('-', '')}">빠른상담 {PHONE}</a></p>
  </section>
</article>""".strip()


def page_shell(row: dict[str, str], body: str) -> str:
    kw = html.escape(row["keyword"])
    canonical = f"/{html.escape(row['slug'])}/"
    return f"""<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <title>{kw} | {BRAND}</title>
  <meta name="description" content="{kw} 작업 범위와 진행 순서, 견적 전 확인사항을 안내합니다.">
  <link rel="canonical" href="{canonical}">
  <style>
    *{{box-sizing:border-box}} body{{margin:0;font-family:Arial,'Noto Sans KR',sans-serif;color:#172033;background:#f6f8fb;line-height:1.75}}
    main{{max-width:920px;margin:auto;background:#fff;min-height:100vh;padding:28px 34px 70px}} h1{{font-size:clamp(28px,5vw,46px);line-height:1.25}}
    h2{{margin-top:42px;font-size:25px}} .eyebrow{{font-weight:700;color:#2459c4}} .hero{{padding:32px 0;border-bottom:1px solid #e5e9f0}}
    li{{margin:9px 0}} .cta{{margin-top:44px;padding:24px;border-radius:16px;background:#edf4ff}} .cta a{{font-size:24px;font-weight:800;color:#164ca4}}
  </style>
</head>
<body><main>{body}</main></body>
</html>"""


def main() -> None:
    p = argparse.ArgumentParser(description="V6.1 키워드별 본문 페이지 생성")
    p.add_argument("--input", default="keywords.xlsx")
    p.add_argument("--output", default="deploys")
    p.add_argument("--limit", type=int, default=0, help="0이면 전체")
    p.add_argument("--overwrite", action="store_true")
    args = p.parse_args()

    input_path, out_dir = Path(args.input), Path(args.output)
    if not input_path.exists():
        raise SystemExit(f"입력 파일 없음: {input_path}")
    rows = load_rows(input_path)
    if args.limit > 0:
        rows = rows[:args.limit]
    out_dir.mkdir(parents=True, exist_ok=True)
    made = skipped = 0
    for i, row in enumerate(rows, 1):
        folder = out_dir / row["slug"]
        target = folder / "index.html"
        if target.exists() and not args.overwrite:
            skipped += 1
            continue
        folder.mkdir(parents=True, exist_ok=True)
        target.write_text(page_shell(row, render_body(row)), encoding="utf-8")
        made += 1
        if i % 500 == 0:
            print(f"진행 {i:,}/{len(rows):,}")
    print(f"[완료] 생성 {made:,}개 / 건너뜀 {skipped:,}개 / 출력 {out_dir.resolve()}")


if __name__ == "__main__":
    main()
