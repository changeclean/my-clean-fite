#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""V6.1 내부링크 생성·삽입 엔진.

기능
- keywords.xlsx 또는 keywords.csv 읽기
- 같은 동네/지역/서비스/공간 유형을 우선해 관련 페이지 선정
- 자기 자신과 중복 링크 제외
- deploys/{slug}/index.html에 관련 청소 정보 영역 삽입 또는 갱신
- link_report_v6_1.csv 결과 저장

실행 예시
    python 4.link_engine_v6_1.py --limit 100
    python 4.link_engine_v6_1.py --max-links 16
    python 4.link_engine_v6_1.py --dry-run
"""
from __future__ import annotations

import argparse
import csv
import html
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl이 필요합니다: pip install openpyxl") from exc

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT = BASE_DIR / "keywords.xlsx"
DEFAULT_OUTPUT = BASE_DIR / "deploys"
DEFAULT_REPORT = BASE_DIR / "link_report_v6_1.csv"

START_MARKER = "<!-- V6_1_INTERNAL_LINKS_START -->"
END_MARKER = "<!-- V6_1_INTERNAL_LINKS_END -->"


@dataclass(frozen=True)
class Page:
    index: int
    keyword: str
    slug: str
    region: str = ""
    neighborhood: str = ""
    service: str = ""
    category: str = ""
    property_type: str = ""
    detail: str = ""


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalized_header(value: Any) -> str:
    return re.sub(r"[\s_-]+", "", clean(value).lower())


HEADER_ALIASES = {
    "keyword": {"keyword", "키워드", "제목", "title"},
    "slug": {"slug", "슬러그", "url", "영문슬러그"},
    "region": {"region", "지역", "지역명", "시군구"},
    "neighborhood": {"neighborhood", "동네", "동", "읍면동", "세부지역"},
    "service": {"service", "서비스", "청소종류", "업종"},
    "category": {"category", "카테고리"},
    "property_type": {"propertytype", "property_type", "공간유형", "건물유형", "주거형태"},
    "detail": {"detail", "상세", "작업내용", "오염", "세부작업"},
}


def canonical_header(header: Any) -> str:
    norm = normalized_header(header)
    for canonical, aliases in HEADER_ALIASES.items():
        if norm in {normalized_header(a) for a in aliases}:
            return canonical
    return clean(header)


def load_xlsx(path: Path) -> list[Page]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration:
        return []
    headers = [canonical_header(v) for v in raw_headers]
    pages: list[Page] = []
    seen: set[str] = set()
    for index, values in enumerate(rows, start=1):
        row = {headers[i]: clean(values[i]) for i in range(min(len(headers), len(values)))}
        keyword, slug = row.get("keyword", ""), row.get("slug", "").strip("/")
        if not keyword or not slug or slug in seen:
            continue
        seen.add(slug)
        pages.append(Page(index=index, keyword=keyword, slug=slug,
                          region=row.get("region", ""), neighborhood=row.get("neighborhood", ""),
                          service=row.get("service", ""), category=row.get("category", ""),
                          property_type=row.get("property_type", ""), detail=row.get("detail", "")))
    return pages


def load_csv(path: Path) -> list[Page]:
    pages: list[Page] = []
    seen: set[str] = set()
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for index, raw in enumerate(reader, start=1):
            row = {canonical_header(k): clean(v) for k, v in raw.items()}
            keyword, slug = row.get("keyword", ""), row.get("slug", "").strip("/")
            if not keyword or not slug or slug in seen:
                continue
            seen.add(slug)
            pages.append(Page(index=index, keyword=keyword, slug=slug,
                              region=row.get("region", ""), neighborhood=row.get("neighborhood", ""),
                              service=row.get("service", ""), category=row.get("category", ""),
                              property_type=row.get("property_type", ""), detail=row.get("detail", "")))
    return pages


def load_pages(path: Path) -> list[Page]:
    if not path.exists():
        raise FileNotFoundError(f"입력 파일이 없습니다: {path}")
    if path.suffix.lower() == ".csv":
        return load_csv(path)
    return load_xlsx(path)


def token_set(text: str) -> set[str]:
    return {x for x in re.findall(r"[가-힣A-Za-z0-9]+", text.lower()) if len(x) >= 2}


def similarity_score(current: Page, other: Page) -> int:
    if current.slug == other.slug:
        return -10_000
    score = 0
    if current.neighborhood and current.neighborhood == other.neighborhood:
        score += 12
    if current.region and current.region == other.region:
        score += 8
    if current.service and current.service == other.service:
        score += 7
    if current.category and current.category == other.category:
        score += 5
    if current.property_type and current.property_type == other.property_type:
        score += 4
    if current.detail and current.detail == other.detail:
        score += 3
    common = token_set(current.keyword) & token_set(other.keyword)
    score += min(len(common), 5)
    # 결과가 늘 같은 앞부분에 몰리지 않도록 안정적인 분산값 추가
    score += sum(ord(ch) for ch in (current.slug + other.slug)) % 3
    return score


def select_links(current: Page, pages: Iterable[Page], max_links: int = 16) -> list[Page]:
    candidates = [(similarity_score(current, p), p) for p in pages if p.slug != current.slug]
    candidates.sort(key=lambda item: (-item[0], item[1].keyword, item[1].slug))
    selected: list[Page] = []
    seen: set[str] = set()
    for score, page in candidates:
        if page.slug in seen:
            continue
        if score <= 0 and selected:
            continue
        selected.append(page)
        seen.add(page.slug)
        if len(selected) >= max_links:
            break
    # 분류 정보가 부족한 엑셀도 링크 수를 채우도록 전체 목록에서 보충
    if len(selected) < max_links:
        for page in pages:
            if page.slug == current.slug or page.slug in seen:
                continue
            selected.append(page)
            seen.add(page.slug)
            if len(selected) >= max_links:
                break
    return selected


def build_internal_links(current: Page, pages: Iterable[Page], max_links: int = 16) -> str:
    selected = select_links(current, pages, max_links=max_links)
    cards = "\n".join(
        f'<li><a href="/{html.escape(p.slug, quote=True)}/">{html.escape(p.keyword)}</a></li>'
        for p in selected
    )
    return f"""{START_MARKER}
<section class="related-cleaning-pages" aria-labelledby="related-cleaning-title">
  <h2 id="related-cleaning-title">함께 확인하면 좋은 청소 정보</h2>
  <ul class="related-cleaning-list">
    {cards}
  </ul>
</section>
<style>
.related-cleaning-pages{{margin-top:42px;padding:24px;border:1px solid #e3e8ef;border-radius:16px;background:#f8fafc}}
.related-cleaning-pages h2{{margin:0 0 16px;font-size:24px}}
.related-cleaning-list{{display:grid;grid-template-columns:repeat(auto-fit,minmax(260px,1fr));gap:10px 18px;margin:0;padding-left:20px}}
.related-cleaning-list li{{margin:0}}
.related-cleaning-list a{{color:#174ea6;text-decoration:none;font-weight:700}}
.related-cleaning-list a:hover{{text-decoration:underline}}
</style>
{END_MARKER}"""


def inject_links(document: str, block: str) -> tuple[str, str]:
    marker_pattern = re.compile(
        re.escape(START_MARKER) + r".*?" + re.escape(END_MARKER),
        flags=re.I | re.S,
    )
    if marker_pattern.search(document):
        return marker_pattern.sub(block, document, count=1), "UPDATED"
    for closing in ("</article>", "</main>", "</body>"):
        match = re.search(re.escape(closing), document, flags=re.I)
        if match:
            pos = match.start()
            return document[:pos] + "\n" + block + "\n" + document[pos:], "INSERTED"
    return document + "\n" + block + "\n", "APPENDED"


def write_report(path: Path, rows: list[dict[str, Any]]) -> None:
    fields = ["index", "keyword", "slug", "status", "link_count", "file", "message"]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows({k: row.get(k, "") for k in fields} for row in rows)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="V6.1 관련 페이지 내부링크 삽입")
    parser.add_argument("--input", default=str(DEFAULT_INPUT), help="keywords.xlsx 또는 keywords.csv")
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT), help="생성 페이지 폴더")
    parser.add_argument("--report", default=str(DEFAULT_REPORT), help="결과 CSV")
    parser.add_argument("--limit", type=int, default=0, help="앞에서 N개만 처리, 0이면 전체")
    parser.add_argument("--max-links", type=int, default=16, help="페이지당 최대 내부링크 수")
    parser.add_argument("--dry-run", action="store_true", help="파일을 저장하지 않고 결과만 확인")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.max_links < 1:
        raise SystemExit("--max-links는 1 이상이어야 합니다.")
    input_path = Path(args.input)
    output_dir = Path(args.output)
    report_path = Path(args.report)
    pages = load_pages(input_path)
    targets = pages[:args.limit] if args.limit > 0 else pages
    results: list[dict[str, Any]] = []

    for number, page in enumerate(targets, start=1):
        target = output_dir / page.slug / "index.html"
        try:
            if not target.exists():
                status, message = "MISSING", "index.html 없음"
                count = 0
            else:
                original = target.read_text(encoding="utf-8")
                links = select_links(page, pages, max_links=args.max_links)
                block = build_internal_links(page, pages, max_links=args.max_links)
                updated, method = inject_links(original, block)
                count = len(links)
                if updated == original:
                    status, message = "SKIP", "변경 없음"
                else:
                    status, message = method, f"내부링크 {count}개"
                    if not args.dry_run:
                        target.write_text(updated, encoding="utf-8")
            results.append({"index": page.index, "keyword": page.keyword, "slug": page.slug,
                            "status": status, "link_count": count, "file": str(target), "message": message})
        except Exception as exc:
            results.append({"index": page.index, "keyword": page.keyword, "slug": page.slug,
                            "status": "ERROR", "link_count": 0, "file": str(target), "message": str(exc)})
        if number % 500 == 0:
            print(f"진행 {number:,}/{len(targets):,}")

    if not args.dry_run:
        write_report(report_path, results)
    success = sum(1 for r in results if r["status"] in {"INSERTED", "UPDATED", "APPENDED", "SKIP"})
    missing = sum(1 for r in results if r["status"] == "MISSING")
    errors = sum(1 for r in results if r["status"] == "ERROR")
    mode = "미리보기" if args.dry_run else "완료"
    print(f"[{mode}] 정상 {success:,} / 파일 없음 {missing:,} / 오류 {errors:,}")
    if not args.dry_run:
        print(f"리포트: {report_path.resolve()}")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
