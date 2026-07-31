#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""V6.1 생성 페이지 품질 검증기.

검사 항목
- deploys/{slug}/index.html 존재 여부
- title / meta description / H1 / canonical
- robots index,follow
- JSON-LD 스키마
- 내부링크 개수 및 중복
- 남은 템플릿 토큰 {{...}}
- 이미지 상대경로 오류
- title / description / canonical 중복
- 페이지별 점수 및 전체 점수

사용 예
    python 5.verify_v6_1.py
    python 5.verify_v6_1.py --limit 100
    python 5.verify_v6_1.py --input keywords.xlsx --deploys deploys
    python 5.verify_v6_1.py --min-links 8 --strict
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple
from urllib.parse import urlparse

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl이 필요합니다: pip install openpyxl") from exc


SITE_URL = "https://changeclean1.netlify.app"
DEFAULT_INPUT = "keywords.xlsx"
DEFAULT_DEPLOYS = "deploys"
DEFAULT_REPORT = "verify_report_v6_1.csv"
DEFAULT_SUMMARY = "verify_summary_v6_1.json"


@dataclass
class KeywordRow:
    index: int
    keyword: str
    slug: str


@dataclass
class VerifyResult:
    index: int
    keyword: str
    slug: str
    path: str
    exists: str
    status: str
    score: float
    title: str
    title_length: int
    description: str
    description_length: int
    h1: str
    canonical: str
    internal_links: int
    duplicate_internal_links: int
    schema_count: int
    token_count: int
    relative_image_errors: int
    issues: str


TITLE_RE = re.compile(r"<title[^>]*>(.*?)</title>", re.I | re.S)
DESCRIPTION_RE = re.compile(
    r'<meta\s+[^>]*name=["\']description["\'][^>]*content=["\'](.*?)["\'][^>]*>',
    re.I | re.S,
)
DESCRIPTION_RE_REVERSED = re.compile(
    r'<meta\s+[^>]*content=["\'](.*?)["\'][^>]*name=["\']description["\'][^>]*>',
    re.I | re.S,
)
H1_RE = re.compile(r"<h1[^>]*>(.*?)</h1>", re.I | re.S)
CANONICAL_RE = re.compile(
    r'<link\s+[^>]*rel=["\']canonical["\'][^>]*href=["\'](.*?)["\'][^>]*>',
    re.I | re.S,
)
CANONICAL_RE_REVERSED = re.compile(
    r'<link\s+[^>]*href=["\'](.*?)["\'][^>]*rel=["\']canonical["\'][^>]*>',
    re.I | re.S,
)
LINK_RE = re.compile(r'<a\s+[^>]*href=["\']([^"\']+)["\']', re.I)
SCHEMA_RE = re.compile(
    r'<script\s+[^>]*type=["\']application/ld\+json["\'][^>]*>(.*?)</script>',
    re.I | re.S,
)
ROBOTS_RE = re.compile(
    r'<meta\s+[^>]*name=["\']robots["\'][^>]*content=["\'](.*?)["\'][^>]*>',
    re.I | re.S,
)
TOKEN_RE = re.compile(r"{{[^{}]+}}")
TAG_RE = re.compile(r"<[^>]+>")
SPACE_RE = re.compile(r"\s+")
RELATIVE_IMAGE_RE = re.compile(r'(?:src|href)=["\']\.\./images/', re.I)


def clean_html_text(value: str) -> str:
    value = TAG_RE.sub(" ", value or "")
    value = value.replace("&nbsp;", " ")
    return SPACE_RE.sub(" ", value).strip()


def first_match(patterns: Iterable[re.Pattern], html: str) -> str:
    for pattern in patterns:
        match = pattern.search(html)
        if match:
            return clean_html_text(match.group(1))
    return ""


def normalize_slug(value: object) -> str:
    text = str(value or "").strip().strip("/")
    return text


def find_column(headers: List[str], candidates: Iterable[str]) -> Optional[int]:
    lowered = {str(h or "").strip().lower(): i for i, h in enumerate(headers)}
    for candidate in candidates:
        if candidate.lower() in lowered:
            return lowered[candidate.lower()]
    return None


def read_keywords(path: Path) -> List[KeywordRow]:
    if not path.exists():
        raise FileNotFoundError(f"입력 파일이 없습니다: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [str(v or "").strip() for v in next(rows)]
    except StopIteration:
        return []

    keyword_col = find_column(headers, ["keyword", "키워드", "title", "제목"])
    slug_col = find_column(headers, ["slug", "슬러그", "url", "path"])

    if keyword_col is None or slug_col is None:
        raise ValueError(
            "keywords.xlsx에서 keyword/키워드 열과 slug/슬러그 열을 찾지 못했습니다."
        )

    results: List[KeywordRow] = []
    for idx, row in enumerate(rows, start=2):
        keyword = str(row[keyword_col] or "").strip() if keyword_col < len(row) else ""
        slug = normalize_slug(row[slug_col] if slug_col < len(row) else "")
        if not keyword or not slug:
            continue
        results.append(KeywordRow(index=idx, keyword=keyword, slug=slug))
    return results


def is_internal_link(href: str, site_url: str) -> bool:
    href = href.strip()
    if not href or href.startswith(("#", "mailto:", "tel:", "javascript:")):
        return False
    if href.startswith("/"):
        return True
    parsed = urlparse(href)
    if not parsed.scheme:
        return True
    site_host = urlparse(site_url).netloc.lower()
    return parsed.netloc.lower() == site_host


def validate_schema(blocks: List[str]) -> Tuple[int, List[str]]:
    valid = 0
    issues: List[str] = []
    for i, raw in enumerate(blocks, start=1):
        try:
            json.loads(raw.strip())
            valid += 1
        except Exception:
            issues.append(f"schema_{i}_json_error")
    return valid, issues


def expected_canonical(site_url: str, slug: str) -> str:
    return f"{site_url.rstrip('/')}/{slug.strip('/')}/"


def verify_page(
    row: KeywordRow,
    deploys_dir: Path,
    site_url: str,
    min_links: int,
) -> VerifyResult:
    page_path = deploys_dir / row.slug / "index.html"
    issues: List[str] = []
    score = 100.0

    if not page_path.exists():
        return VerifyResult(
            index=row.index,
            keyword=row.keyword,
            slug=row.slug,
            path=str(page_path),
            exists="NO",
            status="ERROR",
            score=0.0,
            title="",
            title_length=0,
            description="",
            description_length=0,
            h1="",
            canonical="",
            internal_links=0,
            duplicate_internal_links=0,
            schema_count=0,
            token_count=0,
            relative_image_errors=0,
            issues="missing_file",
        )

    html = page_path.read_text(encoding="utf-8", errors="replace")
    title = first_match([TITLE_RE], html)
    description = first_match([DESCRIPTION_RE, DESCRIPTION_RE_REVERSED], html)
    h1 = first_match([H1_RE], html)
    canonical = first_match([CANONICAL_RE, CANONICAL_RE_REVERSED], html)
    robots = first_match([ROBOTS_RE], html).lower()

    if not title:
        issues.append("missing_title")
        score -= 15
    elif len(title) < 10:
        issues.append("title_too_short")
        score -= 4
    elif len(title) > 65:
        issues.append("title_too_long")
        score -= 3

    if not description:
        issues.append("missing_description")
        score -= 15
    elif len(description) < 45:
        issues.append("description_too_short")
        score -= 4
    elif len(description) > 160:
        issues.append("description_too_long")
        score -= 3

    if not h1:
        issues.append("missing_h1")
        score -= 12
    elif row.keyword not in h1 and h1 not in row.keyword:
        issues.append("h1_keyword_mismatch")
        score -= 3

    expected = expected_canonical(site_url, row.slug)
    if not canonical:
        issues.append("missing_canonical")
        score -= 15
    elif canonical.rstrip("/") != expected.rstrip("/"):
        issues.append("canonical_mismatch")
        score -= 7

    if robots and not ("index" in robots and "follow" in robots):
        issues.append("robots_not_index_follow")
        score -= 5

    schema_blocks = SCHEMA_RE.findall(html)
    valid_schema_count, schema_issues = validate_schema(schema_blocks)
    issues.extend(schema_issues)
    if not schema_blocks:
        issues.append("missing_schema")
        score -= 10
    elif valid_schema_count == 0:
        score -= 8

    hrefs = [href.strip() for href in LINK_RE.findall(html)]
    internal = [href for href in hrefs if is_internal_link(href, site_url)]
    normalized_internal = [href.split("#", 1)[0].rstrip("/") or "/" for href in internal]
    duplicate_internal = len(normalized_internal) - len(set(normalized_internal))

    if len(internal) < min_links:
        issues.append("too_few_internal_links")
        score -= min(10, (min_links - len(internal)) * 1.5)
    if duplicate_internal > 0:
        issues.append("duplicate_internal_links")
        score -= min(5, duplicate_internal * 0.5)

    token_count = len(TOKEN_RE.findall(html))
    if token_count:
        issues.append("remaining_template_tokens")
        score -= min(20, token_count * 4)

    relative_image_errors = len(RELATIVE_IMAGE_RE.findall(html))
    if relative_image_errors:
        issues.append("relative_image_path_error")
        score -= min(10, relative_image_errors * 2)

    if "<html" not in html.lower() or "</html>" not in html.lower():
        issues.append("invalid_html_shell")
        score -= 8

    score = max(0.0, round(score, 1))
    status = "PASS" if score >= 90 and not any(
        issue in issues
        for issue in (
            "missing_title",
            "missing_description",
            "missing_h1",
            "missing_canonical",
            "missing_schema",
            "remaining_template_tokens",
        )
    ) else "CHECK"

    return VerifyResult(
        index=row.index,
        keyword=row.keyword,
        slug=row.slug,
        path=str(page_path),
        exists="YES",
        status=status,
        score=score,
        title=title,
        title_length=len(title),
        description=description,
        description_length=len(description),
        h1=h1,
        canonical=canonical,
        internal_links=len(internal),
        duplicate_internal_links=duplicate_internal,
        schema_count=valid_schema_count,
        token_count=token_count,
        relative_image_errors=relative_image_errors,
        issues="|".join(issues),
    )


def apply_global_duplicate_checks(results: List[VerifyResult]) -> None:
    groups = {
        "duplicate_title": defaultdict(list),
        "duplicate_description": defaultdict(list),
        "duplicate_canonical": defaultdict(list),
    }

    for idx, result in enumerate(results):
        if result.title:
            groups["duplicate_title"][result.title].append(idx)
        if result.description:
            groups["duplicate_description"][result.description].append(idx)
        if result.canonical:
            groups["duplicate_canonical"][result.canonical.rstrip("/")].append(idx)

    penalties = {
        "duplicate_title": 6,
        "duplicate_description": 4,
        "duplicate_canonical": 20,
    }

    for issue, mapping in groups.items():
        for indexes in mapping.values():
            if len(indexes) < 2:
                continue
            for idx in indexes:
                result = results[idx]
                issue_list = [x for x in result.issues.split("|") if x]
                if issue not in issue_list:
                    issue_list.append(issue)
                result.issues = "|".join(issue_list)
                result.score = max(0.0, round(result.score - penalties[issue], 1))
                if issue == "duplicate_canonical" or result.score < 90:
                    result.status = "CHECK"


def write_csv(path: Path, results: List[VerifyResult]) -> None:
    fieldnames = list(asdict(results[0]).keys()) if results else [
        "index", "keyword", "slug", "path", "exists", "status", "score",
        "title", "title_length", "description", "description_length", "h1",
        "canonical", "internal_links", "duplicate_internal_links", "schema_count",
        "token_count", "relative_image_errors", "issues",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for result in results:
            writer.writerow(asdict(result))


def make_summary(results: List[VerifyResult]) -> Dict[str, object]:
    total = len(results)
    pass_count = sum(1 for r in results if r.status == "PASS")
    check_count = sum(1 for r in results if r.status == "CHECK")
    error_count = sum(1 for r in results if r.status == "ERROR")
    average_score = round(sum(r.score for r in results) / total, 1) if total else 0.0

    issue_counter: Counter[str] = Counter()
    for result in results:
        issue_counter.update(x for x in result.issues.split("|") if x)

    return {
        "total_pages": total,
        "pass_pages": pass_count,
        "check_pages": check_count,
        "error_pages": error_count,
        "average_score": average_score,
        "final_score": average_score,
        "issue_counts": dict(issue_counter.most_common()),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="V6.1 생성 페이지 품질 검증기")
    parser.add_argument("--input", default=DEFAULT_INPUT, help="키워드 엑셀 파일")
    parser.add_argument("--deploys", default=DEFAULT_DEPLOYS, help="생성 페이지 폴더")
    parser.add_argument("--report", default=DEFAULT_REPORT, help="CSV 리포트 파일")
    parser.add_argument("--summary", default=DEFAULT_SUMMARY, help="JSON 요약 파일")
    parser.add_argument("--site-url", default=SITE_URL, help="사이트 기준 URL")
    parser.add_argument("--limit", type=int, default=0, help="앞에서 N개만 검사")
    parser.add_argument("--min-links", type=int, default=8, help="최소 내부링크 개수")
    parser.add_argument(
        "--strict",
        action="store_true",
        help="CHECK 또는 ERROR가 있으면 종료코드 1 반환",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = Path(args.input)
    deploys_dir = Path(args.deploys)
    report_path = Path(args.report)
    summary_path = Path(args.summary)

    try:
        rows = read_keywords(input_path)
    except Exception as exc:
        print(f"[오류] {exc}", file=sys.stderr)
        return 2

    if args.limit > 0:
        rows = rows[: args.limit]

    if not rows:
        print("검사할 키워드가 없습니다.")
        return 2

    results: List[VerifyResult] = []
    for i, row in enumerate(rows, start=1):
        results.append(
            verify_page(
                row=row,
                deploys_dir=deploys_dir,
                site_url=args.site_url,
                min_links=max(0, args.min_links),
            )
        )
        if i % 1000 == 0:
            print(f"검사 진행: {i:,}/{len(rows):,}", flush=True)

    apply_global_duplicate_checks(results)
    write_csv(report_path, results)
    summary = make_summary(results)
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print("\n검증 완료")
    print(f"전체 페이지: {summary['total_pages']:,}")
    print(f"PASS: {summary['pass_pages']:,}")
    print(f"CHECK: {summary['check_pages']:,}")
    print(f"ERROR: {summary['error_pages']:,}")
    print(f"최종 점수: {summary['final_score']} / 100")
    print(f"상세 리포트: {report_path}")
    print(f"요약 리포트: {summary_path}")

    if args.strict and (summary["check_pages"] or summary["error_pages"]):
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
