#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""V6.1 원클릭 생성·검증·배포 엔진.

기본 실행:
    python 6.deploy_v6_1.py

100개 로컬 테스트:
    python 6.deploy_v6_1.py --limit 100 --no-git

기존 keywords.xlsx 사용:
    python 6.deploy_v6_1.py --skip-keywords --no-git

전체 강제 재생성 후 Git 배포:
    python 6.deploy_v6_1.py --skip-keywords --force

단계 선택:
    python 6.deploy_v6_1.py --skip-pages --skip-links --no-git
"""
from __future__ import annotations

import argparse
import csv
import html
import os
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Sequence

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit("openpyxl이 필요합니다: pip install openpyxl") from exc


BASE_DIR = Path(__file__).resolve().parent
DEPLOYS_DIR = BASE_DIR / "deploys"
KEYWORDS_FILE = BASE_DIR / "keywords.xlsx"

SITE_URL = "https://changeclean1.netlify.app"
BRAND_NAME = "체인지클린"
SITE_DESCRIPTION = "입주청소, 이사청소, 아파트청소, 원룸청소, 사무실청소, 상가청소 및 폐기물청소 안내"
MAX_URLS_PER_SITEMAP = 45_000

PAGE_ENGINE_CANDIDATES = (
    "2.make_sites_v6_1.py",
    "make_sites_v6_1.py",
)
SEO_ENGINE_CANDIDATES = (
    "3.seo_engine_v6_1.py",
    "seo_engine_v6_1.py",
)
LINK_ENGINE_CANDIDATES = (
    "4.link_engine_v6_1.py",
    "link_engine_v6_1.py",
)
VERIFY_ENGINE_CANDIDATES = (
    "5.verify_v6_1.py",
    "verify_v6_1.py",
)
KEYWORD_ENGINE_CANDIDATES = (
    "1.keyword_engine_v6_1.py",
    "keyword_engine_v6_1.py",
    "keyword_engine_v5.py",
)


@dataclass(frozen=True)
class Page:
    keyword: str
    slug: str
    region: str = ""
    service: str = ""


def log(message: str) -> None:
    print(message, flush=True)


def normalize(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def slugify(text: str) -> str:
    value = normalize(text).lower()
    value = re.sub(r"[^\w가-힣-]+", "-", value, flags=re.UNICODE)
    value = re.sub(r"-{2,}", "-", value).strip("-")
    return value or "page"


def find_script(candidates: Sequence[str], required: bool = True) -> Path | None:
    for name in candidates:
        path = BASE_DIR / name
        if path.is_file():
            return path
    if required:
        raise FileNotFoundError("실행 파일을 찾지 못했습니다: " + ", ".join(candidates))
    return None


def run_command(command: Sequence[str], label: str, allow_nothing_to_commit: bool = False) -> int:
    log("\n" + "=" * 72)
    log(label)
    log("$ " + " ".join(command))
    log("=" * 72)

    try:
        process = subprocess.Popen(
            list(command),
            cwd=BASE_DIR,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            errors="replace",
            bufsize=1,
        )
    except FileNotFoundError:
        log(f"명령을 실행할 수 없습니다: {command[0]}")
        return 127

    assert process.stdout is not None
    output_lines: list[str] = []
    for line in process.stdout:
        print(line, end="")
        output_lines.append(line)

    code = process.wait()
    output = "".join(output_lines).lower()

    if code != 0 and allow_nothing_to_commit and "nothing to commit" in output:
        return 0
    if code != 0:
        log(f"\n오류: {label} 단계가 종료 코드 {code}로 실패했습니다.")
    return code


def read_pages(limit: int = 0) -> list[Page]:
    if not KEYWORDS_FILE.exists():
        raise FileNotFoundError(f"키워드 파일이 없습니다: {KEYWORDS_FILE}")

    workbook = load_workbook(KEYWORDS_FILE, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    try:
        header_row = next(rows)
    except StopIteration:
        workbook.close()
        return []

    headers = [normalize(v).lower() for v in header_row]
    aliases = {
        "keyword": ("keyword", "키워드", "키워드명", "title"),
        "slug": ("slug", "슬러그", "url", "url_slug"),
        "region": ("region", "지역", "지역명", "city"),
        "service": ("service", "서비스", "청소종류", "category", "카테고리"),
    }

    def column_index(keys: Sequence[str]) -> int | None:
        for key in keys:
            key_lower = key.lower()
            if key_lower in headers:
                return headers.index(key_lower)
        return None

    keyword_idx = column_index(aliases["keyword"])
    slug_idx = column_index(aliases["slug"])
    region_idx = column_index(aliases["region"])
    service_idx = column_index(aliases["service"])

    if keyword_idx is None:
        workbook.close()
        raise ValueError("keywords.xlsx에서 keyword 또는 키워드 열을 찾지 못했습니다.")

    pages: list[Page] = []
    seen_slugs: set[str] = set()

    for values in rows:
        keyword = normalize(values[keyword_idx] if keyword_idx < len(values) else "")
        if not keyword:
            continue

        raw_slug = normalize(values[slug_idx] if slug_idx is not None and slug_idx < len(values) else "")
        slug = slugify(raw_slug or keyword)
        if slug in seen_slugs:
            continue
        seen_slugs.add(slug)

        region = normalize(values[region_idx] if region_idx is not None and region_idx < len(values) else "")
        service = normalize(values[service_idx] if service_idx is not None and service_idx < len(values) else "")
        pages.append(Page(keyword=keyword, slug=slug, region=region, service=service))

        if limit and len(pages) >= limit:
            break

    workbook.close()
    return pages


def xml_escape(value: str) -> str:
    return html.escape(value, quote=True)


def write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8", newline="\n")


def build_urlset(urls: Iterable[str], priority: str = "0.8") -> str:
    today = datetime.now().strftime("%Y-%m-%d")
    items = []
    for url in urls:
        items.append(
            "  <url>\n"
            f"    <loc>{xml_escape(url)}</loc>\n"
            f"    <lastmod>{today}</lastmod>\n"
            "    <changefreq>weekly</changefreq>\n"
            f"    <priority>{priority}</priority>\n"
            "  </url>"
        )
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n'
        + "\n".join(items)
        + "\n</urlset>\n"
    )


def build_sitemap_index(filenames: Iterable[str]) -> str:
    today = datetime.now().strftime("%Y-%m-%d")
    items = []
    for filename in filenames:
        url = f"{SITE_URL.rstrip('/')}/{filename}"
        items.append(
            "  <sitemap>\n"
            f"    <loc>{xml_escape(url)}</loc>\n"
            f"    <lastmod>{today}</lastmod>\n"
            "  </sitemap>"
        )
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<sitemapindex xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n'
        + "\n".join(items)
        + "\n</sitemapindex>\n"
    )


def generate_sitemaps(pages: Sequence[Page]) -> list[str]:
    DEPLOYS_DIR.mkdir(parents=True, exist_ok=True)

    main_urls = [f"{SITE_URL.rstrip('/')}/"]
    page_urls = [f"{SITE_URL.rstrip('/')}/{page.slug}/" for page in pages]

    sitemap_files = ["sitemap-main.xml"]
    write_text(DEPLOYS_DIR / "sitemap-main.xml", build_urlset(main_urls, priority="1.0"))

    for start in range(0, len(page_urls), MAX_URLS_PER_SITEMAP):
        number = start // MAX_URLS_PER_SITEMAP + 1
        filename = f"sitemap-pages-{number}.xml"
        write_text(
            DEPLOYS_DIR / filename,
            build_urlset(page_urls[start : start + MAX_URLS_PER_SITEMAP], priority="0.8"),
        )
        sitemap_files.append(filename)

    write_text(DEPLOYS_DIR / "sitemap.xml", build_sitemap_index(sitemap_files))
    return sitemap_files


def generate_robots() -> None:
    write_text(
        DEPLOYS_DIR / "robots.txt",
        "User-agent: *\n"
        "Allow: /\n\n"
        f"Sitemap: {SITE_URL.rstrip('/')}/sitemap.xml\n",
    )


def generate_rss(pages: Sequence[Page], max_items: int = 100) -> None:
    now = datetime.now(timezone.utc).strftime("%a, %d %b %Y %H:%M:%S +0000")
    items = []

    for page in list(pages)[-max_items:][::-1]:
        url = f"{SITE_URL.rstrip('/')}/{page.slug}/"
        items.append(
            "    <item>\n"
            f"      <title>{xml_escape(page.keyword)}</title>\n"
            f"      <link>{xml_escape(url)}</link>\n"
            f"      <guid isPermaLink=\"true\">{xml_escape(url)}</guid>\n"
            f"      <description>{xml_escape(page.keyword + ' 청소 서비스 안내')}</description>\n"
            f"      <pubDate>{now}</pubDate>\n"
            "    </item>"
        )

    rss = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<rss version="2.0">\n'
        "  <channel>\n"
        f"    <title>{xml_escape(BRAND_NAME)}</title>\n"
        f"    <link>{xml_escape(SITE_URL)}</link>\n"
        f"    <description>{xml_escape(SITE_DESCRIPTION)}</description>\n"
        "    <language>ko</language>\n"
        f"    <lastBuildDate>{now}</lastBuildDate>\n"
        + "\n".join(items)
        + "\n  </channel>\n"
        "</rss>\n"
    )
    write_text(DEPLOYS_DIR / "rss.xml", rss)


def validate_xml_files(files: Iterable[Path]) -> None:
    for path in files:
        try:
            ET.parse(path)
        except ET.ParseError as exc:
            raise RuntimeError(f"XML 문법 오류: {path.name}: {exc}") from exc


def count_generated_pages() -> int:
    if not DEPLOYS_DIR.exists():
        return 0
    count = 0
    for index_file in DEPLOYS_DIR.glob("*/index.html"):
        if index_file.parent.name not in {"region", "regions", "service", "services"}:
            count += 1
    return count


def write_deploy_report(pages: Sequence[Page], sitemap_files: Sequence[str], status: str) -> None:
    report_path = BASE_DIR / "deploy_report_v6_1.csv"
    with report_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=[
                "generated_at",
                "status",
                "keyword_rows",
                "generated_pages",
                "sitemap_files",
                "site_url",
            ],
        )
        writer.writeheader()
        writer.writerow(
            {
                "generated_at": datetime.now().isoformat(timespec="seconds"),
                "status": status,
                "keyword_rows": len(pages),
                "generated_pages": count_generated_pages(),
                "sitemap_files": len(sitemap_files),
                "site_url": SITE_URL,
            }
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="V6.1 원클릭 생성·검증·배포 엔진")
    parser.add_argument("--limit", type=int, default=0, help="테스트할 키워드/페이지 수")
    parser.add_argument("--force", action="store_true", help="페이지 전체 강제 재생성")
    parser.add_argument("--skip-keywords", action="store_true", help="키워드 생성 단계 건너뛰기")
    parser.add_argument("--skip-pages", action="store_true", help="페이지 생성 단계 건너뛰기")
    parser.add_argument("--skip-seo", action="store_true", help="별도 SEO 엔진 단계 건너뛰기")
    parser.add_argument("--skip-links", action="store_true", help="내부링크 단계 건너뛰기")
    parser.add_argument("--skip-verify", action="store_true", help="검증 단계 건너뛰기")
    parser.add_argument("--no-git", action="store_true", help="Git add/commit/push 건너뛰기")
    parser.add_argument("--dry-run", action="store_true", help="Git 명령만 실제 실행하지 않기")
    parser.add_argument("--message", default="", help="Git commit 메시지")
    parser.add_argument("--max-links", type=int, default=12, help="페이지당 내부링크 수")
    parser.add_argument("--allow-verify-fail", action="store_true", help="검증 실패여도 Git 단계 진행")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    python = sys.executable
    DEPLOYS_DIR.mkdir(parents=True, exist_ok=True)

    log("V6.1 배포 엔진 시작")
    log(f"작업 폴더: {BASE_DIR}")
    log(f"사이트: {SITE_URL}")

    # 1. 키워드 생성
    if not args.skip_keywords:
        keyword_engine = find_script(KEYWORD_ENGINE_CANDIDATES, required=False)
        if keyword_engine:
            command = [python, keyword_engine.name]
            if keyword_engine.name == "keyword_engine_v5.py":
                command.append("--append")
            if args.limit:
                command += ["--limit", str(args.limit)]
            code = run_command(command, "1단계: 키워드 생성")
            if code:
                return code
        elif KEYWORDS_FILE.exists():
            log("\n1단계: 키워드 엔진이 없어 기존 keywords.xlsx를 사용합니다.")
        else:
            log("\n오류: 키워드 엔진과 keywords.xlsx가 모두 없습니다.")
            return 1
    else:
        log("\n1단계: 키워드 생성 건너뜀")

    if not KEYWORDS_FILE.exists():
        log(f"오류: {KEYWORDS_FILE.name} 파일이 없습니다.")
        return 1

    # 2. 페이지 생성
    if not args.skip_pages:
        try:
            page_engine = find_script(PAGE_ENGINE_CANDIDATES)
        except FileNotFoundError as exc:
            log(str(exc))
            return 1

        command = [python, page_engine.name]
        if args.limit:
            command += ["--limit", str(args.limit)]
        if args.force:
            command.append("--force")

        code = run_command(command, "2단계: V6.1 페이지 생성")
        if code:
            return code
    else:
        log("\n2단계: 페이지 생성 건너뜀")

    # 3. SEO 엔진
    # make_sites가 SEO 모듈을 직접 import하는 구조일 수 있어, CLI 지원 여부를 확인해 선택 실행합니다.
    if not args.skip_seo:
        seo_engine = find_script(SEO_ENGINE_CANDIDATES, required=False)
        if seo_engine:
            help_result = subprocess.run(
                [python, seo_engine.name, "--help"],
                cwd=BASE_DIR,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
            )
            help_text = (help_result.stdout or "") + (help_result.stderr or "")
            if "--all" in help_text or "--input" in help_text:
                command = [python, seo_engine.name]
                if "--all" in help_text:
                    command.append("--all")
                if args.limit and "--limit" in help_text:
                    command += ["--limit", str(args.limit)]
                code = run_command(command, "3단계: V6.1 SEO 적용")
                if code:
                    return code
            else:
                log("\n3단계: SEO 엔진은 모듈형입니다. 페이지 생성기가 사용한 것으로 처리합니다.")
        else:
            log("\n3단계: SEO 엔진 파일이 없어 건너뜁니다.")
    else:
        log("\n3단계: SEO 적용 건너뜀")

    # 4. 내부링크
    if not args.skip_links:
        try:
            link_engine = find_script(LINK_ENGINE_CANDIDATES)
        except FileNotFoundError as exc:
            log(str(exc))
            return 1

        command = [python, link_engine.name, "--max-links", str(max(1, args.max_links))]
        if args.limit:
            command += ["--limit", str(args.limit)]

        code = run_command(command, "4단계: 내부링크 적용")
        if code:
            return code
    else:
        log("\n4단계: 내부링크 적용 건너뜀")

    # 5. 사이트맵/RSS/robots
    try:
        pages = read_pages(limit=args.limit)
    except Exception as exc:
        log(f"\n키워드 읽기 오류: {exc}")
        return 1

    if not pages:
        log("\n오류: 사이트맵에 넣을 페이지가 없습니다.")
        return 1

    log("\n" + "=" * 72)
    log("5단계: sitemap.xml · RSS · robots.txt 생성")
    log("=" * 72)

    sitemap_files = generate_sitemaps(pages)
    generate_rss(pages)
    generate_robots()

    xml_paths = [DEPLOYS_DIR / "sitemap.xml", DEPLOYS_DIR / "rss.xml"]
    xml_paths.extend(DEPLOYS_DIR / name for name in sitemap_files)
    try:
        validate_xml_files(xml_paths)
    except RuntimeError as exc:
        log(str(exc))
        return 1

    log(f"키워드 기준 URL: {len(pages):,}개")
    log(f"생성된 페이지 폴더: {count_generated_pages():,}개")
    log("생성 파일:")
    log("- deploys/sitemap.xml")
    for name in sitemap_files:
        log(f"- deploys/{name}")
    log("- deploys/rss.xml")
    log("- deploys/robots.txt")

    # 6. 검증
    verify_code = 0
    if not args.skip_verify:
        try:
            verify_engine = find_script(VERIFY_ENGINE_CANDIDATES)
        except FileNotFoundError as exc:
            log(str(exc))
            return 1

        command = [python, verify_engine.name]
        if args.limit:
            command += ["--limit", str(args.limit)]

        verify_code = run_command(command, "6단계: V6.1 최종 검증")
        if verify_code and not args.allow_verify_fail:
            write_deploy_report(pages, sitemap_files, "VERIFY_FAILED")
            log("\n검증 실패로 Git 배포를 중단했습니다.")
            log("문제 수정 후 다시 실행하거나 --allow-verify-fail을 사용하세요.")
            return verify_code
    else:
        log("\n6단계: 최종 검증 건너뜀")

    write_deploy_report(
        pages,
        sitemap_files,
        "VERIFY_WARNING" if verify_code else "SUCCESS",
    )

    # 7. Git 배포
    if args.no_git:
        log("\n7단계: Git 배포 건너뜀")
    else:
        message = args.message or f"V6.1 deploy {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        commands = (
            ["git", "add", "."],
            ["git", "commit", "-m", message],
            ["git", "push"],
        )
        for command in commands:
            if args.dry_run:
                log("\nDRY RUN: " + " ".join(command))
                continue
            code = run_command(
                command,
                "7단계: Git 배포 - " + " ".join(command[:2]),
                allow_nothing_to_commit=(command[1] == "commit"),
            )
            if code:
                return code

    log("\n" + "=" * 72)
    log("V6.1 전체 작업 완료")
    log("=" * 72)
    log(SITE_URL)
    log(f"{SITE_URL}/sitemap.xml")
    log(f"{SITE_URL}/rss.xml")
    log("리포트: deploy_report_v6_1.csv")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
