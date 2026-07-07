"""
keyword_builder.py

청소업체 SEO 키워드 자동 생성기

기능:
- 서울/인천/경기 주요 지역 + 주력지역(인천, 김포, 부천, 고양, 일산) 우선 반영
- 입주청소/이사청소/거주청소/준공청소/상가청소/사무실청소 등 키워드 생성
- 번호 없는 SEO 슬러그 생성
- 중복 키워드/슬러그 자동 제거
- keywords.xlsx 생성

실행:
python keyword_builder.py

생성 결과:
keywords.xlsx
"""

from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


OUTPUT_FILE = "keywords.xlsx"


# =========================
# 지역 DB
# 필요하면 여기에 동 이름을 계속 추가하면 됩니다.
# =========================

AREAS = [
    # 주력 지역: 인천
    ("인천", "미추홀구", "incheon-michuhol-gu"),
    ("인천", "숭의동", "incheon-sungui-dong"),
    ("인천", "주안동", "incheon-juan-dong"),
    ("인천", "용현동", "incheon-yonghyeon-dong"),
    ("인천", "도화동", "incheon-dohwa-dong"),
    ("인천", "학익동", "incheon-hagik-dong"),
    ("인천", "연수구", "incheon-yeonsu-gu"),
    ("인천", "송도동", "incheon-songdo-dong"),
    ("인천", "청학동", "incheon-cheonghak-dong"),
    ("인천", "동춘동", "incheon-dongchun-dong"),
    ("인천", "남동구", "incheon-namdong-gu"),
    ("인천", "구월동", "incheon-guwol-dong"),
    ("인천", "간석동", "incheon-ganseok-dong"),
    ("인천", "논현동", "incheon-nonhyeon-dong"),
    ("인천", "만수동", "incheon-mansu-dong"),
    ("인천", "부평구", "incheon-bupyeong-gu"),
    ("인천", "부평동", "incheon-bupyeong-dong"),
    ("인천", "부개동", "incheon-bugae-dong"),
    ("인천", "삼산동", "incheon-samsan-dong"),
    ("인천", "청천동", "incheon-cheongcheon-dong"),
    ("인천", "계양구", "incheon-gyeyang-gu"),
    ("인천", "계산동", "incheon-gyesan-dong"),
    ("인천", "작전동", "incheon-jakjeon-dong"),
    ("인천", "서구", "incheon-seo-gu"),
    ("인천", "청라동", "incheon-cheongna-dong"),
    ("인천", "검단동", "incheon-geomdan-dong"),
    ("인천", "마전동", "incheon-majeon-dong"),
    ("인천", "당하동", "incheon-dangha-dong"),
    ("인천", "가정동", "incheon-gajeong-dong"),
    ("인천", "중구", "incheon-jung-gu"),
    ("인천", "운서동", "incheon-unseo-dong"),
    ("인천", "영종동", "incheon-yeongjong-dong"),
    ("인천", "동구", "incheon-dong-gu"),
    ("인천", "송림동", "incheon-songnim-dong"),

    # 주력 지역: 김포
    ("경기", "김포시", "gyeonggi-gimpo-si"),
    ("경기", "장기동", "gyeonggi-gimpo-janggi-dong"),
    ("경기", "운양동", "gyeonggi-gimpo-unyang-dong"),
    ("경기", "구래동", "gyeonggi-gimpo-gurae-dong"),
    ("경기", "마산동", "gyeonggi-gimpo-masan-dong"),
    ("경기", "풍무동", "gyeonggi-gimpo-pungmu-dong"),
    ("경기", "사우동", "gyeonggi-gimpo-sau-dong"),
    ("경기", "고촌읍", "gyeonggi-gimpo-gochon-eup"),
    ("경기", "양촌읍", "gyeonggi-gimpo-yangchon-eup"),
    ("경기", "통진읍", "gyeonggi-gimpo-tongjin-eup"),

    # 주력 지역: 부천
    ("경기", "부천시", "gyeonggi-bucheon-si"),
    ("경기", "상동", "gyeonggi-bucheon-sang-dong"),
    ("경기", "중동", "gyeonggi-bucheon-jung-dong"),
    ("경기", "심곡동", "gyeonggi-bucheon-simgok-dong"),
    ("경기", "송내동", "gyeonggi-bucheon-songnae-dong"),
    ("경기", "역곡동", "gyeonggi-bucheon-yeokgok-dong"),
    ("경기", "소사동", "gyeonggi-bucheon-sosa-dong"),
    ("경기", "괴안동", "gyeonggi-bucheon-goean-dong"),
    ("경기", "옥길동", "gyeonggi-bucheon-okgil-dong"),

    # 주력 지역: 고양/일산
    ("경기", "고양시", "gyeonggi-goyang-si"),
    ("경기", "고양시 덕양구", "gyeonggi-goyang-deogyang-gu"),
    ("경기", "화정동", "gyeonggi-goyang-hwajeong-dong"),
    ("경기", "행신동", "gyeonggi-goyang-haengsin-dong"),
    ("경기", "원흥동", "gyeonggi-goyang-wonheung-dong"),
    ("경기", "삼송동", "gyeonggi-goyang-samsong-dong"),
    ("경기", "고양시 일산동구", "gyeonggi-goyang-ilsandong-gu"),
    ("경기", "마두동", "gyeonggi-ilsan-madu-dong"),
    ("경기", "백석동", "gyeonggi-ilsan-baekseok-dong"),
    ("경기", "장항동", "gyeonggi-ilsan-janghang-dong"),
    ("경기", "풍동", "gyeonggi-ilsan-pung-dong"),
    ("경기", "고양시 일산서구", "gyeonggi-goyang-ilsanseo-gu"),
    ("경기", "주엽동", "gyeonggi-ilsan-juyeop-dong"),
    ("경기", "탄현동", "gyeonggi-ilsan-tanhyeon-dong"),
    ("경기", "대화동", "gyeonggi-ilsan-daehwa-dong"),

    # 서울 25개 구
    ("서울", "강남구", "seoul-gangnam-gu"),
    ("서울", "강동구", "seoul-gangdong-gu"),
    ("서울", "강북구", "seoul-gangbuk-gu"),
    ("서울", "강서구", "seoul-gangseo-gu"),
    ("서울", "관악구", "seoul-gwanak-gu"),
    ("서울", "광진구", "seoul-gwangjin-gu"),
    ("서울", "구로구", "seoul-guro-gu"),
    ("서울", "금천구", "seoul-geumcheon-gu"),
    ("서울", "노원구", "seoul-nowon-gu"),
    ("서울", "도봉구", "seoul-dobong-gu"),
    ("서울", "동대문구", "seoul-dongdaemun-gu"),
    ("서울", "동작구", "seoul-dongjak-gu"),
    ("서울", "마포구", "seoul-mapo-gu"),
    ("서울", "서대문구", "seoul-seodaemun-gu"),
    ("서울", "서초구", "seoul-seocho-gu"),
    ("서울", "성동구", "seoul-seongdong-gu"),
    ("서울", "성북구", "seoul-seongbuk-gu"),
    ("서울", "송파구", "seoul-songpa-gu"),
    ("서울", "양천구", "seoul-yangcheon-gu"),
    ("서울", "영등포구", "seoul-yeongdeungpo-gu"),
    ("서울", "용산구", "seoul-yongsan-gu"),
    ("서울", "은평구", "seoul-eunpyeong-gu"),
    ("서울", "종로구", "seoul-jongno-gu"),
    ("서울", "중구", "seoul-jung-gu"),
    ("서울", "중랑구", "seoul-jungnang-gu"),

    # 경기 주요 시군
    ("경기", "수원시", "gyeonggi-suwon-si"),
    ("경기", "성남시", "gyeonggi-seongnam-si"),
    ("경기", "용인시", "gyeonggi-yongin-si"),
    ("경기", "화성시", "gyeonggi-hwaseong-si"),
    ("경기", "안산시", "gyeonggi-ansan-si"),
    ("경기", "안양시", "gyeonggi-anyang-si"),
    ("경기", "평택시", "gyeonggi-pyeongtaek-si"),
    ("경기", "시흥시", "gyeonggi-siheung-si"),
    ("경기", "광명시", "gyeonggi-gwangmyeong-si"),
    ("경기", "군포시", "gyeonggi-gunpo-si"),
    ("경기", "의왕시", "gyeonggi-uiwang-si"),
    ("경기", "하남시", "gyeonggi-hanam-si"),
    ("경기", "광주시", "gyeonggi-gwangju-si"),
    ("경기", "오산시", "gyeonggi-osan-si"),
    ("경기", "이천시", "gyeonggi-icheon-si"),
    ("경기", "안성시", "gyeonggi-anseong-si"),
    ("경기", "의정부시", "gyeonggi-uijeongbu-si"),
    ("경기", "양주시", "gyeonggi-yangju-si"),
    ("경기", "파주시", "gyeonggi-paju-si"),
    ("경기", "남양주시", "gyeonggi-namyangju-si"),
    ("경기", "구리시", "gyeonggi-guri-si"),
    ("경기", "포천시", "gyeonggi-pocheon-si"),
    ("경기", "양평군", "gyeonggi-yangpyeong-gun"),
    ("경기", "가평군", "gyeonggi-gapyeong-gun"),
    ("경기", "여주시", "gyeonggi-yeoju-si"),
]


SERVICES = [
    ("입주청소", "move-in-cleaning"),
    ("이사청소", "moving-cleaning"),
    ("거주청소", "living-cleaning"),
    ("준공청소", "post-construction-cleaning"),
    ("상가청소", "commercial-cleaning"),
    ("사무실청소", "office-cleaning"),
]


BUILDINGS = [
    ("", ""),
    ("아파트", "apartment"),
    ("빌라", "villa"),
    ("원룸", "studio"),
    ("오피스텔", "officetel"),
    ("신축", "new-home"),
    ("구축", "old-home"),
    ("24평", "24pyeong"),
    ("34평", "34pyeong"),
    ("40평", "40pyeong"),
]


INTENTS = [
    ("", ""),
    ("추천", "recommended"),
    ("전문업체", "professional"),
    ("잘하는 곳", "best"),
    ("비용", "cost"),
    ("가격", "price"),
    ("후기", "review"),
    ("당일", "same-day"),
    ("곰팡이 제거", "mold-removal"),
    ("새집증후군", "new-house-syndrome"),
]


def make_keyword(province: str, area: str, building: str, service: str, intent: str) -> str:
    parts = [province, area]
    if building:
        parts.append(building)
    parts.append(service)
    if intent:
        parts.append(intent)
    return " ".join(parts)


def make_slug(area_slug: str, building_slug: str, service_slug: str, intent_slug: str) -> str:
    parts = [area_slug]
    if building_slug:
        parts.append(building_slug)
    parts.append(service_slug)
    if intent_slug:
        parts.append(intent_slug)
    return "-".join(parts)


def build_keywords(limit: int | None = None):
    rows = []
    seen_keywords = set()
    seen_slugs = set()

    for province, area, area_slug in AREAS:
        for service, service_slug in SERVICES:
            for building, building_slug in BUILDINGS:
                for intent, intent_slug in INTENTS:
                    keyword = make_keyword(province, area, building, service, intent)
                    slug = make_slug(area_slug, building_slug, service_slug, intent_slug)

                    if keyword in seen_keywords or slug in seen_slugs:
                        continue

                    seen_keywords.add(keyword)
                    seen_slugs.add(slug)
                    rows.append((keyword, slug))

                    if limit and len(rows) >= limit:
                        return rows

    return rows


def save_xlsx(rows, output_file: str = OUTPUT_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "keywords"
    ws.append(["키워드", "SEO슬러그"])

    for keyword, slug in rows:
        ws.append([keyword, slug])

    ws.column_dimensions[get_column_letter(1)].width = 55
    ws.column_dimensions[get_column_letter(2)].width = 80

    wb.save(output_file)


def main():
    rows = build_keywords()
    save_xlsx(rows)

    print("=" * 50)
    print("청소업체 SEO 키워드 생성 완료")
    print("=" * 50)
    print(f"생성 개수: {len(rows):,}개")
    print(f"저장 파일: {Path(OUTPUT_FILE).resolve()}")
    print()
    print("다음 실행:")
    print("python make_sites_v4.py")


if __name__ == "__main__":
    main()
