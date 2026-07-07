"""
keyword_builder_v5.py

청소업체 SEO 키워드 자동 생성기 v5

실행:
python keyword_builder_v5.py

기존 생성기에 바로 적용:
python keyword_builder_v5.py --overwrite

테스트용 개수 제한:
python keyword_builder_v5.py --limit 5000
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import List, Tuple

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


OUTPUT_FILE = "keywords_v5.xlsx"
OVERWRITE_FILE = "keywords.xlsx"


AREAS = [
    ("인천", "숭의동", "incheon-sungui-dong"),
    ("인천", "주안동", "incheon-juan-dong"),
    ("인천", "용현동", "incheon-yonghyeon-dong"),
    ("인천", "도화동", "incheon-dohwa-dong"),
    ("인천", "학익동", "incheon-hagik-dong"),
    ("인천", "송도동", "incheon-songdo-dong"),
    ("인천", "구월동", "incheon-guwol-dong"),
    ("인천", "간석동", "incheon-ganseok-dong"),
    ("인천", "논현동", "incheon-nonhyeon-dong"),
    ("인천", "만수동", "incheon-mansu-dong"),
    ("인천", "부평동", "incheon-bupyeong-dong"),
    ("인천", "청라동", "incheon-cheongna-dong"),
    ("인천", "검단동", "incheon-geomdan-dong"),
    ("인천", "영종동", "incheon-yeongjong-dong"),
    ("경기", "김포시", "gyeonggi-gimpo-si"),
    ("경기", "장기동", "gyeonggi-gimpo-janggi-dong"),
    ("경기", "운양동", "gyeonggi-gimpo-unyang-dong"),
    ("경기", "구래동", "gyeonggi-gimpo-gurae-dong"),
    ("경기", "풍무동", "gyeonggi-gimpo-pungmu-dong"),
    ("경기", "부천시", "gyeonggi-bucheon-si"),
    ("경기", "상동", "gyeonggi-bucheon-sang-dong"),
    ("경기", "중동", "gyeonggi-bucheon-jung-dong"),
    ("경기", "송내동", "gyeonggi-bucheon-songnae-dong"),
    ("경기", "고양시", "gyeonggi-goyang-si"),
    ("경기", "화정동", "gyeonggi-goyang-hwajeong-dong"),
    ("경기", "행신동", "gyeonggi-goyang-haengsin-dong"),
    ("경기", "일산동구", "gyeonggi-goyang-ilsandong-gu"),
    ("경기", "일산서구", "gyeonggi-goyang-ilsanseo-gu"),
    ("경기", "마두동", "gyeonggi-ilsan-madu-dong"),
    ("경기", "백석동", "gyeonggi-ilsan-baekseok-dong"),
    ("경기", "주엽동", "gyeonggi-ilsan-juyeop-dong"),
    ("서울", "강남구", "seoul-gangnam-gu"),
    ("서울", "강동구", "seoul-gangdong-gu"),
    ("서울", "강북구", "seoul-gangbuk-gu"),
    ("서울", "강서구", "seoul-gangseo-gu"),
    ("서울", "관악구", "seoul-gwanak-gu"),
    ("서울", "광진구", "seoul-gwangjin-gu"),
    ("서울", "구로구", "seoul-guro-gu"),
    ("서울", "금천구", "seoul-geumcheon-gu"),
    ("서울", "노원구", "seoul-nowon-gu"),
    ("서울", "도봉구", "seoul-dobong-gu"),
    ("서울", "동대문구", "seoul-dongdaemun-gu"),
    ("서울", "동작구", "seoul-dongjak-gu"),
    ("서울", "마포구", "seoul-mapo-gu"),
    ("서울", "서대문구", "seoul-seodaemun-gu"),
    ("서울", "서초구", "seoul-seocho-gu"),
    ("서울", "성동구", "seoul-seongdong-gu"),
    ("서울", "성북구", "seoul-seongbuk-gu"),
    ("서울", "송파구", "seoul-songpa-gu"),
    ("서울", "양천구", "seoul-yangcheon-gu"),
    ("서울", "영등포구", "seoul-yeongdeungpo-gu"),
    ("서울", "용산구", "seoul-yongsan-gu"),
    ("서울", "은평구", "seoul-eunpyeong-gu"),
    ("서울", "종로구", "seoul-jongno-gu"),
    ("서울", "중구", "seoul-jung-gu"),
    ("서울", "중랑구", "seoul-jungnang-gu"),
    ("경기", "수원시", "gyeonggi-suwon-si"),
    ("경기", "성남시", "gyeonggi-seongnam-si"),
    ("경기", "용인시", "gyeonggi-yongin-si"),
    ("경기", "화성시", "gyeonggi-hwaseong-si"),
    ("경기", "안산시", "gyeonggi-ansan-si"),
    ("경기", "안양시", "gyeonggi-anyang-si"),
    ("경기", "평택시", "gyeonggi-pyeongtaek-si"),
    ("경기", "시흥시", "gyeonggi-siheung-si"),
    ("경기", "광명시", "gyeonggi-gwangmyeong-si"),
    ("경기", "군포시", "gyeonggi-gunpo-si"),
    ("경기", "의왕시", "gyeonggi-uiwang-si"),
    ("경기", "하남시", "gyeonggi-hanam-si"),
    ("경기", "광주시", "gyeonggi-gwangju-si"),
    ("경기", "오산시", "gyeonggi-osan-si"),
    ("경기", "파주시", "gyeonggi-paju-si"),
    ("경기", "남양주시", "gyeonggi-namyangju-si"),
]


SERVICES = [
    ("입주청소", "move-in-cleaning"),
    ("이사청소", "moving-cleaning"),
    ("거주청소", "living-cleaning"),
    ("준공청소", "post-construction-cleaning"),
    ("상가청소", "commercial-cleaning"),
    ("사무실청소", "office-cleaning"),
    ("원룸청소", "studio-cleaning"),
    ("오피스텔청소", "officetel-cleaning"),
    ("빌라청소", "villa-cleaning"),
    ("아파트청소", "apartment-cleaning"),
    ("새집증후군청소", "new-house-syndrome-cleaning"),
    ("곰팡이청소", "mold-cleaning"),
    ("유리창청소", "window-cleaning"),
    ("화장실청소", "bathroom-cleaning"),
    ("주방청소", "kitchen-cleaning"),
]


BUILDINGS = [
    ("", ""),
    ("아파트", "apartment"),
    ("빌라", "villa"),
    ("원룸", "studio"),
    ("오피스텔", "officetel"),
    ("신축", "new-home"),
    ("구축", "old-home"),
    ("상가", "commercial-space"),
    ("사무실", "office-space"),
    ("24평", "24pyeong"),
    ("30평", "30pyeong"),
    ("34평", "34pyeong"),
    ("40평", "40pyeong"),
]


INTENTS = [
    ("", ""),
    ("추천", "recommended"),
    ("전문업체", "professional"),
    ("잘하는 곳", "best"),
    ("비용", "cost"),
    ("가격", "price"),
    ("견적", "quote"),
    ("후기", "review"),
    ("당일", "same-day"),
    ("예약", "reservation"),
    ("상담", "consultation"),
    ("저렴한 곳", "affordable"),
    ("꼼꼼한 곳", "thorough"),
    ("업체 비교", "company-comparison"),
]


PATTERNS = [
    ("{province} {area} {service}", "{area_slug}-{service_slug}"),
    ("{province} {area} {service} {intent}", "{area_slug}-{service_slug}-{intent_slug}"),
    ("{province} {area} {building} {service}", "{area_slug}-{building_slug}-{service_slug}"),
    ("{province} {area} {building} {service} {intent}", "{area_slug}-{building_slug}-{service_slug}-{intent_slug}"),
    ("{province} {area} {service} 잘하는 청소업체", "{area_slug}-{service_slug}-cleaning-company-best"),
    ("{province} {area} {service} 믿을만한 업체", "{area_slug}-{service_slug}-trusted-company"),
    ("{province} {area} {service} 빠른 상담", "{area_slug}-{service_slug}-quick-consultation"),
    ("{province} {area} {service} 예약 가능한 곳", "{area_slug}-{service_slug}-reservation-available"),
    ("{province} {area} {service} 전후 비교", "{area_slug}-{service_slug}-before-after"),
    ("{province} {area} {service} 현장 후기", "{area_slug}-{service_slug}-site-review"),
    ("{province} {area} {building} 청소 견적", "{area_slug}-{building_slug}-cleaning-quote"),
    ("{province} {area} {building} 청소 비용", "{area_slug}-{building_slug}-cleaning-cost"),
    ("{province} {area} {building} 청소 추천", "{area_slug}-{building_slug}-cleaning-recommended"),
    ("{province} {area} {building} 청소 전문업체", "{area_slug}-{building_slug}-cleaning-professional"),
]


def clean_spaces(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def clean_slug(slug: str) -> str:
    slug = slug.replace("--", "-")
    slug = re.sub(r"-+", "-", slug)
    return slug.strip("-")


def build_keywords(limit: int | None = None) -> List[Tuple[str, str]]:
    rows: List[Tuple[str, str]] = []
    seen_keywords = set()
    seen_slugs = set()

    for province, area, area_slug in AREAS:
        for service, service_slug in SERVICES:
            for building, building_slug in BUILDINGS:
                for intent, intent_slug in INTENTS:
                    c = {
                        "province": province,
                        "area": area,
                        "area_slug": area_slug,
                        "service": service,
                        "service_slug": service_slug,
                        "building": building,
                        "building_slug": building_slug,
                        "intent": intent,
                        "intent_slug": intent_slug,
                    }

                    for kp, sp in PATTERNS:
                        if "{building}" in kp and not building:
                            continue
                        if "{intent}" in kp and not intent:
                            continue
                        if "{building_slug}" in sp and not building_slug:
                            continue
                        if "{intent_slug}" in sp and not intent_slug:
                            continue

                        keyword = clean_spaces(kp.format(**c))
                        slug = clean_slug(sp.format(**c))

                        if keyword in seen_keywords or slug in seen_slugs:
                            continue

                        seen_keywords.add(keyword)
                        seen_slugs.add(slug)
                        rows.append((keyword, slug))

                        if limit and len(rows) >= limit:
                            return rows

    return rows


def save_xlsx(rows: List[Tuple[str, str]], output_file: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "keywords"
    ws.append(["키워드", "SEO슬러그"])

    for keyword, slug in rows:
        ws.append([keyword, slug])

    ws.column_dimensions[get_column_letter(1)].width = 60
    ws.column_dimensions[get_column_letter(2)].width = 95
    wb.save(output_file)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="청소업체 SEO 키워드 자동 생성기 v5")
    parser.add_argument("--limit", type=int, default=None, help="생성 개수 제한")
    parser.add_argument("--overwrite", action="store_true", help="keywords.xlsx로 저장해서 기존 생성기에 바로 사용")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    rows = build_keywords(limit=args.limit)
    output = OVERWRITE_FILE if args.overwrite else OUTPUT_FILE
    save_xlsx(rows, output)

    print("=" * 60)
    print("청소업체 SEO 키워드 v5 생성 완료")
    print("=" * 60)
    print(f"생성 개수: {len(rows):,}개")
    print(f"저장 파일: {Path(output).resolve()}")
    print()
    if args.overwrite:
        print("다음 실행:")
        print("python make_sites_v4_1.py")
    else:
        print("기존 생성기에 바로 쓰려면 파일명을 keywords.xlsx로 바꾸거나")
        print("python keyword_builder_v5.py --overwrite 를 실행하세요.")


if __name__ == "__main__":
    main()
