"""
make_sites_v4_1.py

청소업체 SEO CMS v4.1 - 속도개선 버전

v4 대비 개선:
- 내부링크 인덱스 사전 생성
- 페이지마다 전체 7만개를 매번 비교하지 않음
- 진행률에 속도/예상 남은 시간 표시
- 7만개 이상 대량 생성에 적합

필요 파일:
- keywords.xlsx
- my_template/index.html
- keyword_parser_v4.py
- seo_engine_v4.py
- content_engine_v4.py
- link_engine_v4_1.py
- sitemap_engine_v4.py
- quality_engine_v4.py

실행:
python make_sites_v4_1.py

변경된 페이지만 생성:
python make_sites_v4_1.py --incremental

Git 자동 Push:
python make_sites_v4_1.py --git-push
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import shutil
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import load_workbook

from keyword_parser_v4 import parse_keyword
from seo_engine_v4 import build_seo_pack
from content_engine_v4 import build_content_pack
from link_engine_v4_1 import build_link_index, build_link_pack_fast, make_page_record
from sitemap_engine_v4 import generate_all_sitemaps
from quality_engine_v4 import inspect_html, write_quality_report


BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_DIR = BASE_DIR / "my_template"
TEMPLATE_FILE = TEMPLATE_DIR / "index.html"
OUTPUT_DIR = BASE_DIR / "deploys"
CACHE_FILE = BASE_DIR / ".seo_cms_v4_cache.csv"

BASE_URL = "https://changeclean.co.kr/city"

KEYWORD_FILES = [
    BASE_DIR / "keywords.xlsx",
    BASE_DIR / "SEO_키워드_고유슬러그_샘플.xlsx",
    BASE_DIR / "SEO_롱테일키워드_고유슬러그.xlsx",
]

STATIC_DIRS = [
    BASE_DIR / "images",
    TEMPLATE_DIR / "images",
    TEMPLATE_DIR / "assets",
]


def log(msg: str) -> None:
    print(msg, flush=True)


def find_keyword_file() -> Path:
    for f in KEYWORD_FILES:
        if f.exists():
            return f
    raise FileNotFoundError("keywords.xlsx 파일을 프로젝트 폴더에 넣어주세요.")


def ensure_template() -> None:
    if not TEMPLATE_FILE.exists():
        raise FileNotFoundError("my_template/index.html 파일이 없습니다.")


def read_keywords(path: Path) -> List[Tuple[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    rows = []
    seen = set()

    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not row or len(row) < 2:
            continue

        keyword = str(row[0] or "").strip()
        slug = str(row[1] or "").strip().strip("/")

        if idx == 1 and ("키워드" in keyword or "슬러그" in slug or "slug" in slug.lower()):
            continue

        if not keyword or not slug:
            continue

        if slug in seen:
            continue

        seen.add(slug)
        rows.append((keyword, slug))

    if not rows:
        raise ValueError("엑셀에서 키워드/슬러그를 읽지 못했습니다.")

    return rows


def page_hash(keyword: str, slug: str) -> str:
    return hashlib.sha1(f"{keyword}|{slug}".encode("utf-8")).hexdigest()


def load_cache() -> Dict[str, str]:
    if not CACHE_FILE.exists():
        return {}

    out = {}
    with CACHE_FILE.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            if row.get("slug"):
                out[row["slug"]] = row.get("hash", "")
    return out


def save_cache(pages: List[Dict[str, str]]) -> None:
    with CACHE_FILE.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["slug", "hash", "updated_at"])
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        for p in pages:
            writer.writerow([p["slug"], p["hash"], now])


def prepare_output(incremental: bool) -> None:
    if not incremental and OUTPUT_DIR.exists():
        shutil.rmtree(OUTPUT_DIR)
    OUTPUT_DIR.mkdir(exist_ok=True)

    for src in STATIC_DIRS:
        if not src.exists() or not src.is_dir():
            continue

        dst = OUTPUT_DIR / src.name

        if dst.exists() and not incremental:
            shutil.rmtree(dst)

        if not dst.exists():
            shutil.copytree(src, dst)
            log(f"📁 복사: {src.name}")


def read_template() -> str:
    return TEMPLATE_FILE.read_text(encoding="utf-8")


def fix_relative_paths(html: str) -> str:
    fixes = {
        'href="assets/': 'href="../assets/',
        "href='assets/": "href='../assets/",
        'src="assets/': 'src="../assets/',
        "src='assets/": "src='../assets/",
        'href="images/': 'href="../images/',
        "href='images/": "href='../images/",
        'src="images/': 'src="../images/',
        "src='images/": "src='../images/",
    }
    for a, b in fixes.items():
        html = html.replace(a, b)
    return html


def inject_head(html: str, head: str) -> str:
    if "{{SEO_HEAD}}" in html:
        return html.replace("{{SEO_HEAD}}", head)
    if "</head>" in html:
        return html.replace("</head>", head + "\n</head>")
    return head + "\n" + html


def replace_all(html: str, repl: Dict[str, str]) -> str:
    for k, v in repl.items():
        html = html.replace(k, str(v))
    return html


def remove_leftover_placeholders(html: str) -> str:
    import re
    return re.sub(r"\{\{[^{}]+\}\}", "", html)


def write_page(slug: str, html: str) -> None:
    out = OUTPUT_DIR / slug
    out.mkdir(parents=True, exist_ok=True)
    (out / "index.html").write_text(html, encoding="utf-8")


def build_pages(rows: List[Tuple[str, str]]) -> List[Dict[str, str]]:
    pages = []
    for keyword, slug in rows:
        parsed = parse_keyword(keyword)
        page = make_page_record(keyword, slug, parsed)
        page["hash"] = page_hash(keyword, slug)
        pages.append(page)
    return pages


def write_generation_log(pages: List[Dict[str, str]]) -> None:
    with (OUTPUT_DIR / "generation_log.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["keyword", "slug", "area", "service", "building", "intent", "pyeong"])
        for p in pages:
            writer.writerow([
                p.get("keyword", ""),
                p.get("slug", ""),
                p.get("area", ""),
                p.get("service", ""),
                p.get("building", ""),
                p.get("intent", ""),
                p.get("pyeong", ""),
            ])


def format_eta(seconds: float) -> str:
    seconds = max(0, int(seconds))
    h = seconds // 3600
    m = (seconds % 3600) // 60
    s = seconds % 60
    if h:
        return f"{h}시간 {m}분 {s}초"
    if m:
        return f"{m}분 {s}초"
    return f"{s}초"


def git_push() -> None:
    msg = f"SEO CMS v4.1 update {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    cmds = [
        ["git", "add", "."],
        ["git", "commit", "-m", msg],
        ["git", "push"],
    ]
    for cmd in cmds:
        result = subprocess.run(cmd, cwd=BASE_DIR, text=True, capture_output=True)
        if result.returncode != 0:
            all_text = (result.stdout + result.stderr).lower()
            if "nothing to commit" in all_text:
                log("변경사항 없음")
                return
            log(result.stdout)
            log(result.stderr)
            raise RuntimeError("git 명령 실패")
    log("✅ Git push 완료")


def run(incremental: bool = False, do_git_push: bool = False) -> None:
    log("🚀 청소업체 SEO CMS v4.1 속도개선 버전 시작")
    ensure_template()

    keyword_file = find_keyword_file()
    log(f"📄 키워드 파일: {keyword_file.name}")

    rows = read_keywords(keyword_file)
    pages = build_pages(rows)
    total = len(pages)
    log(f"✅ 페이지 데이터: {total:,}개")

    log("🔗 내부링크 인덱스 생성 중...")
    link_index = build_link_index(pages)
    log("✅ 내부링크 인덱스 완료")

    prepare_output(incremental)
    template = read_template()
    cache = load_cache() if incremental else {}

    quality_rows = []
    generated = 0
    skipped = 0
    start_time = time.time()

    for i, page in enumerate(pages):
        slug = page["slug"]

        if incremental and cache.get(slug) == page["hash"] and (OUTPUT_DIR / slug / "index.html").exists():
            skipped += 1
            continue

        html = template

        seo = build_seo_pack(page, BASE_URL)
        content = build_content_pack(page, pages)
        links = build_link_pack_fast(page, pages, link_index, BASE_URL)

        html = inject_head(html, seo["{{SEO_HEAD}}"])

        repl = {}
        repl.update(seo)
        repl.update(content)
        repl.update(links)
        repl.update({
            "{{키워드}}": page.get("keyword", ""),
            "{{지역명}}": page.get("area", ""),
            "{{서비스}}": page.get("service", ""),
            "{{건물}}": page.get("building", ""),
            "{{의도}}": page.get("intent", ""),
            "{{평수}}": page.get("pyeong", ""),
            "{{URL}}": f"{BASE_URL}/{slug}/",
            "{{슬러그}}": slug,
        })

        html = replace_all(html, repl)
        html = fix_relative_paths(html)
        html = remove_leftover_placeholders(html)

        write_page(slug, html)
        quality_rows.append(inspect_html(slug, html))
        generated += 1

        if generated % 500 == 0 or i + 1 == total:
            elapsed = time.time() - start_time
            speed = generated / elapsed if elapsed > 0 else 0
            remaining_count = total - (i + 1)
            eta = remaining_count / speed if speed > 0 else 0
            log(
                f"[{i+1:,}/{total:,}] 생성:{generated:,} 건너뜀:{skipped:,} "
                f"속도:{speed:.1f}p/s 남은시간:{format_eta(eta)}"
            )

    write_generation_log(pages)
    write_quality_report(OUTPUT_DIR / "quality_report.csv", quality_rows)
    save_cache(pages)
    generate_all_sitemaps(OUTPUT_DIR, pages, BASE_URL)

    log("")
    log("✅ v4.1 생성 완료")
    log(f"생성: {generated:,}개 / 건너뜀: {skipped:,}개")
    log(f"결과: {OUTPUT_DIR}")

    if do_git_push:
        git_push()


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--incremental", action="store_true", help="변경된 페이지만 생성")
    parser.add_argument("--git-push", action="store_true", help="생성 후 git push")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    try:
        run(incremental=args.incremental, do_git_push=args.git_push)
    except Exception as e:
        log("❌ 오류 발생")
        log(str(e))
        sys.exit(1)
