#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
V6.1 통합 원클릭 엔진
====================

키워드 생성은 하지 않습니다.
기존 keywords.xlsx를 읽어서 아래 작업을 한 번에 처리합니다.

1. 신규 페이지만 생성
2. SEO 메타태그 및 JSON-LD 적용
3. 내부링크 자동 적용
4. 메인 페이지 생성
5. sitemap.xml / RSS / robots.txt 생성
6. 전체 품질검사
7. 선택적으로 Git add / commit / push

권장 테스트:
    python v6_1_all_in_one.py --limit 100 --no-git

전체 실행:
    python v6_1_all_in_one.py

기존 페이지를 다시 생성:
    python v6_1_all_in_one.py --force

Git 없이 전체 생성:
    python v6_1_all_in_one.py --no-git
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import re
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Sequence
from urllib.parse import urlparse

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit("openpyxl이 필요합니다. 먼저 실행: pip install openpyxl") from exc


# ============================================================
# 기본 설정
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
KEYWORDS_FILE = BASE_DIR / "keywords.xlsx"
DEPLOYS_DIR = BASE_DIR / "deploys"

SITE_URL = "https://clever-griffin-93819d.netlify.app"
BRAND_NAME = "체인지클린"
PHONE = "1688-6751"

DEFAULT_MAX_LINKS = 12
MAX_URLS_PER_SITEMAP = 45_000
RSS_MAX_ITEMS = 100

VERIFY_REPORT = BASE_DIR / "verify_report_v6_1.csv"
VERIFY_SUMMARY = BASE_DIR / "verify_summary_v6_1.json"
DEPLOY_REPORT = BASE_DIR / "deploy_report_v6_1.csv"

LINK_START = "<!-- V6_1_INTERNAL_LINKS_START -->"
LINK_END = "<!-- V6_1_INTERNAL_LINKS_END -->"


# ============================================================
# 데이터 구조
# ============================================================

@dataclass
class Page:
    keyword: str
    slug: str
    region: str = ""
    neighborhood: str = ""
    service: str = ""
    category: str = ""
    property_type: str = ""
    detail: str = ""

    @property
    def url(self) -> str:
        return f"{SITE_URL.rstrip('/')}/{self.slug.strip('/')}/"

    @property
    def path(self) -> Path:
        return DEPLOYS_DIR / self.slug / "index.html"


@dataclass
class VerifyResult:
    keyword: str
    slug: str
    status: str
    score: int
    issues: str
    title: str
    description: str
    canonical: str
    h1: str
    internal_links: int


# ============================================================
# 공통 유틸
# ============================================================

def log(message: str = "") -> None:
    print(message, flush=True)


def clean(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def slugify(value: str) -> str:
    text = clean(value).lower()
    text = re.sub(r"[^\w가-힣-]+", "-", text, flags=re.UNICODE)
    text = re.sub(r"-{2,}", "-", text).strip("-")
    return text or "page"


def esc(value: str) -> str:
    return html.escape(value, quote=True)


def deterministic_choice(seed: str, items: Sequence[str]) -> str:
    digest = hashlib.sha256(seed.encode("utf-8")).hexdigest()
    return items[int(digest[:8], 16) % len(items)]


def write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8", newline="\n")


def strip_tags(value: str) -> str:
    return re.sub(r"<[^>]+>", " ", value)


def normalize_compare(value: str) -> str:
    return re.sub(r"\s+", " ", html.unescape(strip_tags(value))).strip().lower()


# ============================================================
# keywords.xlsx 읽기
# ============================================================

COLUMN_ALIASES = {
    "keyword": ("keyword", "키워드", "키워드명", "title", "제목"),
    "slug": ("slug", "슬러그", "url", "url_slug", "영문슬러그"),
    "region": ("region", "지역", "지역명", "city", "시군구"),
    "neighborhood": ("neighborhood", "동", "동명", "읍면동", "area"),
    "service": ("service", "서비스", "청소종류", "업종"),
    "category": ("category", "카테고리", "분류"),
    "property_type": ("property_type", "property", "건물유형", "주거유형", "공간유형"),
    "detail": ("detail", "세부", "상세", "작업내용", "포인트"),
}


def read_pages(limit: int = 0) -> list[Page]:
    if not KEYWORDS_FILE.exists():
        raise FileNotFoundError(f"keywords.xlsx가 없습니다: {KEYWORDS_FILE}")

    workbook = load_workbook(KEYWORDS_FILE, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    try:
        raw_headers = next(rows)
    except StopIteration:
        workbook.close()
        return []

    headers = [clean(v).lower() for v in raw_headers]

    def find_index(field: str) -> int | None:
        for alias in COLUMN_ALIASES[field]:
            alias_lower = alias.lower()
            if alias_lower in headers:
                return headers.index(alias_lower)
        return None

    indexes = {field: find_index(field) for field in COLUMN_ALIASES}
    if indexes["keyword"] is None:
        workbook.close()
        raise ValueError("keyword 또는 키워드 열을 찾지 못했습니다.")

    pages: list[Page] = []
    seen_slugs: set[str] = set()
    seen_keywords: set[str] = set()

    for row in rows:
        def get(field: str) -> str:
            idx = indexes[field]
            if idx is None or idx >= len(row):
                return ""
            return clean(row[idx])

        keyword = get("keyword")
        if not keyword:
            continue

        keyword_key = keyword.lower()
        raw_slug = get("slug")
        slug = slugify(raw_slug or keyword)

        if keyword_key in seen_keywords or slug in seen_slugs:
            continue

        seen_keywords.add(keyword_key)
        seen_slugs.add(slug)

        pages.append(
            Page(
                keyword=keyword,
                slug=slug,
                region=get("region"),
                neighborhood=get("neighborhood"),
                service=get("service"),
                category=get("category"),
                property_type=get("property_type"),
                detail=get("detail"),
            )
        )

        if limit and len(pages) >= limit:
            break

    workbook.close()
    return pages


# ============================================================
# 페이지 내용 및 SEO 생성
# ============================================================

def infer_service(page: Page) -> str:
    if page.service:
        return page.service

    services = (
        "입주청소", "이사청소", "상가청소", "사무실청소",
        "폐기물청소", "원룸청소", "아파트청소", "거주청소",
    )
    for service in services:
        if service in page.keyword:
            return service
    return "전문청소"


def infer_region(page: Page) -> str:
    if page.neighborhood:
        return page.neighborhood
    if page.region:
        return page.region
    parts = page.keyword.split()
    return " ".join(parts[:2]) if parts else "지역"


def make_title(page: Page) -> str:
    title = page.keyword.strip()
    suffix = deterministic_choice(
        page.slug + "title",
        (" | 체인지클린", " - 청소 작업 안내", " | 현장별 청소 정보"),
    )
    max_keyword = 60 - len(suffix)
    return title[:max_keyword].rstrip() + suffix


def make_description(page: Page) -> str:
    region = infer_region(page)
    service = infer_service(page)
    variants = (
        f"{page.keyword} 작업 범위와 준비사항, 진행 순서, 확인 항목을 정리했습니다. {region} {service} 현장에 맞춰 필요한 청소 내용을 확인해보세요.",
        f"{page.keyword} 관련 청소 과정과 체크포인트를 안내합니다. 공간 상태와 오염 범위를 확인한 뒤 현장에 맞는 작업을 진행합니다.",
        f"{region} {service}를 알아보는 분을 위한 안내입니다. {page.keyword} 작업 전 확인사항부터 마무리 점검까지 자세히 정리했습니다.",
    )
    description = deterministic_choice(page.slug + "desc", variants)
    return description[:155].rstrip()


def make_json_ld(page: Page, title: str, description: str) -> str:
    service = infer_service(page)
    region = infer_region(page)

    data = {
        "@context": "https://schema.org",
        "@graph": [
            {
                "@type": "LocalBusiness",
                "@id": f"{SITE_URL.rstrip('/')}/#business",
                "name": BRAND_NAME,
                "url": SITE_URL,
                "telephone": PHONE,
                "areaServed": {
                    "@type": "AdministrativeArea",
                    "name": region,
                },
            },
            {
                "@type": "Service",
                "@id": f"{page.url}#service",
                "name": page.keyword,
                "serviceType": service,
                "description": description,
                "url": page.url,
                "provider": {
                    "@id": f"{SITE_URL.rstrip('/')}/#business",
                },
                "areaServed": {
                    "@type": "AdministrativeArea",
                    "name": region,
                },
            },
            {
                "@type": "BreadcrumbList",
                "@id": f"{page.url}#breadcrumb",
                "itemListElement": [
                    {
                        "@type": "ListItem",
                        "position": 1,
                        "name": "홈",
                        "item": SITE_URL.rstrip("/") + "/",
                    },
                    {
                        "@type": "ListItem",
                        "position": 2,
                        "name": title,
                        "item": page.url,
                    },
                ],
            },
        ],
    }
    return json.dumps(data, ensure_ascii=False, separators=(",", ":"))


def make_content(page: Page) -> str:
    region = infer_region(page)
    service = infer_service(page)
    property_type = page.property_type or "현장"
    detail = page.detail or "오염 상태와 공간 구조"

    intro_variants = (
        f"{page.keyword} 작업은 단순히 눈에 보이는 먼지만 제거하는 과정이 아닙니다. {property_type}의 구조와 {detail}을 함께 확인해야 작업 누락을 줄일 수 있습니다.",
        f"{page.keyword}를 진행할 때는 공간별 오염도와 사용 흔적을 먼저 살펴야 합니다. 같은 면적이라도 현장 상태에 따라 필요한 장비와 작업 순서가 달라질 수 있습니다.",
        f"{region}에서 {service}를 준비한다면 바닥, 창틀, 주방, 욕실 등 구역별 상태를 먼저 확인하는 것이 좋습니다. 현장 조건에 맞춘 순서로 진행해야 마무리 품질을 높일 수 있습니다.",
    )
    intro = deterministic_choice(page.slug + "intro", intro_variants)

    return f"""
<section class="intro">
  <p>{esc(intro)}</p>
  <p>{esc(region)} {esc(service)}는 공간의 크기뿐 아니라 오염 범위, 짐의 유무, 작업 난이도를 함께 확인해 진행합니다.</p>
</section>

<section>
  <h2>{esc(page.keyword)} 작업 전 확인사항</h2>
  <ul>
    <li>현장 면적과 방·욕실 개수 확인</li>
    <li>창틀, 유리, 주방 및 욕실 오염 상태 확인</li>
    <li>가구와 폐기물 유무 확인</li>
    <li>주차와 엘리베이터 사용 가능 여부 확인</li>
    <li>집중 작업이 필요한 부분 사전 확인</li>
  </ul>
</section>

<section>
  <h2>청소 진행 순서</h2>
  <ol>
    <li>전체 공간과 오염 상태를 확인합니다.</li>
    <li>천장과 벽면, 수납장 등 위쪽부터 먼지를 제거합니다.</li>
    <li>창틀, 주방, 욕실 등 오염이 많은 구역을 세척합니다.</li>
    <li>바닥 청소 후 잔여 오염과 물기를 확인합니다.</li>
    <li>구역별 마무리 점검을 진행합니다.</li>
  </ol>
</section>

<section>
  <h2>구역별 체크포인트</h2>
  <h3>주방</h3>
  <p>상·하부장 내부와 외부, 후드, 싱크대, 가스레인지 주변의 기름때와 분진을 확인합니다.</p>
  <h3>욕실</h3>
  <p>수전, 거울, 배수구, 타일, 환풍구 주변의 물때와 먼지를 확인합니다.</p>
  <h3>창틀과 유리</h3>
  <p>창틀 틈새의 모래와 분진을 제거하고 유리 표면의 얼룩을 정리합니다.</p>
  <h3>바닥</h3>
  <p>공간 전체의 먼지와 잔여 오염을 제거한 뒤 모서리와 걸레받이를 확인합니다.</p>
</section>

<section>
  <h2>작업 후 확인할 부분</h2>
  <p>청소가 끝난 뒤에는 눈높이에서 보이지 않는 수납장 내부, 창틀 모서리, 배수구 주변을 함께 확인하는 것이 좋습니다. 현장 상태에 따라 제거가 어려운 변색이나 손상은 오염과 구분하여 확인해야 합니다.</p>
</section>

<section class="contact">
  <h2>{esc(region)} {esc(service)} 상담</h2>
  <p>현장 사진과 면적, 필요한 작업 범위를 알려주시면 상담에 도움이 됩니다.</p>
  <p><strong>빠른상담: {esc(PHONE)}</strong></p>
</section>
""".strip()


def make_page_html(page: Page) -> str:
    title = make_title(page)
    description = make_description(page)
    json_ld = make_json_ld(page, title, description)
    content = make_content(page)

    return f"""<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{esc(title)}</title>
  <meta name="description" content="{esc(description)}">
  <meta name="robots" content="index,follow,max-image-preview:large">
  <link rel="canonical" href="{esc(page.url)}">

  <meta property="og:type" content="article">
  <meta property="og:locale" content="ko_KR">
  <meta property="og:site_name" content="{esc(BRAND_NAME)}">
  <meta property="og:title" content="{esc(title)}">
  <meta property="og:description" content="{esc(description)}">
  <meta property="og:url" content="{esc(page.url)}">

  <meta name="twitter:card" content="summary">
  <meta name="twitter:title" content="{esc(title)}">
  <meta name="twitter:description" content="{esc(description)}">

  <script type="application/ld+json">{json_ld}</script>

  <style>
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      color: #222;
      background: #f6f8fa;
      font-family: Arial, "Noto Sans KR", sans-serif;
      line-height: 1.75;
    }}
    header {{
      background: #fff;
      border-bottom: 1px solid #e5e7eb;
    }}
    .wrap {{
      width: min(960px, calc(100% - 32px));
      margin: 0 auto;
    }}
    header .wrap {{
      padding: 18px 0;
      font-size: 20px;
      font-weight: 700;
    }}
    main {{
      margin: 30px auto;
      padding: 34px;
      background: #fff;
      border-radius: 16px;
      box-shadow: 0 8px 28px rgba(0,0,0,.06);
    }}
    h1 {{ margin-top: 0; line-height: 1.35; font-size: 32px; }}
    h2 {{ margin-top: 38px; font-size: 24px; }}
    h3 {{ margin-top: 24px; font-size: 19px; }}
    ul, ol {{ padding-left: 24px; }}
    a {{ color: #1259a7; text-decoration: none; }}
    a:hover {{ text-decoration: underline; }}
    .intro {{
      padding: 22px;
      background: #f3f7fb;
      border-radius: 12px;
    }}
    .contact {{
      padding: 22px;
      margin-top: 38px;
      background: #fff7e6;
      border-radius: 12px;
    }}
    .related-cleaning-pages {{
      margin-top: 42px;
      padding-top: 26px;
      border-top: 1px solid #e5e7eb;
    }}
    .related-cleaning-pages ul {{
      columns: 2;
    }}
    footer {{
      padding: 24px 0 50px;
      text-align: center;
      color: #667085;
    }}
    @media (max-width: 640px) {{
      main {{ padding: 22px; }}
      h1 {{ font-size: 27px; }}
      .related-cleaning-pages ul {{ columns: 1; }}
    }}
  </style>
</head>
<body>
<header>
  <div class="wrap"><a href="/">{esc(BRAND_NAME)}</a></div>
</header>

<main class="wrap">
  <article>
    <h1>{esc(page.keyword)}</h1>
    {content}
  </article>
</main>

<footer>
  <div class="wrap">© {datetime.now().year} {esc(BRAND_NAME)}</div>
</footer>
</body>
</html>
"""


# ============================================================
# 페이지 생성
# ============================================================

def generate_pages(pages: Sequence[Page], force: bool = False) -> tuple[int, int]:
    created = 0
    skipped = 0

    for index, page in enumerate(pages, start=1):
        if page.path.exists() and not force:
            skipped += 1
        else:
            write_text(page.path, make_page_html(page))
            created += 1

        if index % 500 == 0 or index == len(pages):
            log(f"페이지 처리: {index:,}/{len(pages):,}")

    return created, skipped


# ============================================================
# 내부링크
# ============================================================

TOKEN_RE = re.compile(r"[가-힣A-Za-z0-9]+")


def page_tokens(page: Page) -> set[str]:
    return {
        token.lower()
        for token in TOKEN_RE.findall(page.keyword)
        if len(token) >= 2
    }


def related_score(source: Page, target: Page, source_tokens: set[str], target_tokens: set[str]) -> int:
    if source.slug == target.slug:
        return -1

    score = 0
    if source.neighborhood and source.neighborhood == target.neighborhood:
        score += 100
    if source.region and source.region == target.region:
        score += 60
    if source.service and source.service == target.service:
        score += 50
    if source.category and source.category == target.category:
        score += 35
    if source.property_type and source.property_type == target.property_type:
        score += 25
    if source.detail and source.detail == target.detail:
        score += 15

    score += len(source_tokens & target_tokens) * 8
    return score


def build_related_html(source: Page, targets: Sequence[Page]) -> str:
    links = "\n".join(
        f'    <li><a href="/{esc(target.slug)}/">{esc(target.keyword)}</a></li>'
        for target in targets
    )
    return f"""{LINK_START}
<section class="related-cleaning-pages">
  <h2>함께 확인할 청소 정보</h2>
  <ul>
{links}
  </ul>
</section>
{LINK_END}"""


def inject_related_links(content: str, block: str) -> str:
    pattern = re.compile(
        re.escape(LINK_START) + r".*?" + re.escape(LINK_END),
        flags=re.DOTALL,
    )
    if pattern.search(content):
        return pattern.sub(block, content)

    for marker in ("</article>", "</main>", "</body>"):
        if marker in content:
            return content.replace(marker, block + "\n  " + marker, 1)

    return content + "\n" + block


def apply_internal_links(pages: Sequence[Page], max_links: int) -> int:
    tokens = {page.slug: page_tokens(page) for page in pages}
    updated = 0

    # 대량 사이트에서 모든 페이지끼리 완전 비교하면 매우 느리므로
    # 지역/서비스/키워드 토큰 인덱스로 후보군을 좁힙니다.
    by_region: dict[str, list[Page]] = {}
    by_service: dict[str, list[Page]] = {}
    by_token: dict[str, list[Page]] = {}

    for page in pages:
        if page.region:
            by_region.setdefault(page.region, []).append(page)
        if page.service:
            by_service.setdefault(page.service, []).append(page)
        for token in list(tokens[page.slug])[:8]:
            by_token.setdefault(token, []).append(page)

    for index, source in enumerate(pages, start=1):
        candidates: dict[str, Page] = {}

        for candidate in by_region.get(source.region, []):
            candidates[candidate.slug] = candidate
        for candidate in by_service.get(source.service, []):
            candidates[candidate.slug] = candidate
        for token in list(tokens[source.slug])[:8]:
            for candidate in by_token.get(token, []):
                candidates[candidate.slug] = candidate

        # 후보가 적으면 인접 페이지를 보충
        if len(candidates) < max_links * 2:
            position = index - 1
            start = max(0, position - 30)
            end = min(len(pages), position + 31)
            for candidate in pages[start:end]:
                candidates[candidate.slug] = candidate

        ranked = sorted(
            (
                (
                    related_score(
                        source,
                        target,
                        tokens[source.slug],
                        tokens[target.slug],
                    ),
                    target,
                )
                for target in candidates.values()
                if target.slug != source.slug
            ),
            key=lambda item: (-item[0], item[1].slug),
        )

        related = [target for score, target in ranked if score >= 0][:max_links]
        if not related:
            continue

        if not source.path.exists():
            continue

        content = source.path.read_text(encoding="utf-8", errors="replace")
        new_content = inject_related_links(content, build_related_html(source, related))

        if new_content != content:
            write_text(source.path, new_content)
            updated += 1

        if index % 500 == 0 or index == len(pages):
            log(f"내부링크 처리: {index:,}/{len(pages):,}")

    return updated


# ============================================================
# 메인 페이지
# ============================================================

def generate_home(pages: Sequence[Page]) -> None:
    sample_pages = list(pages[:60])
    links = "\n".join(
        f'<li><a href="/{esc(page.slug)}/">{esc(page.keyword)}</a></li>'
        for page in sample_pages
    )

    content = f"""<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{esc(BRAND_NAME)} | 입주·이사·상가·사무실 청소</title>
  <meta name="description" content="지역과 공간 유형별 입주청소, 이사청소, 아파트청소, 원룸청소, 사무실청소, 상가청소 정보를 확인하세요.">
  <meta name="robots" content="index,follow,max-image-preview:large">
  <link rel="canonical" href="{esc(SITE_URL.rstrip('/') + '/')}">
  <style>
    body {{ margin:0; font-family:Arial,"Noto Sans KR",sans-serif; color:#222; background:#f6f8fa; line-height:1.7; }}
    main {{ width:min(1000px,calc(100% - 32px)); margin:32px auto; padding:36px; background:white; border-radius:16px; }}
    h1 {{ font-size:36px; }}
    ul {{ columns:2; padding-left:22px; }}
    a {{ color:#1259a7; text-decoration:none; }}
    @media(max-width:640px) {{ main{{padding:22px}} ul{{columns:1}} }}
  </style>
</head>
<body>
<main>
  <h1>{esc(BRAND_NAME)} 청소 정보</h1>
  <p>지역과 현장 조건에 맞는 청소 작업 범위와 준비사항을 확인할 수 있습니다.</p>
  <p><strong>빠른상담: {esc(PHONE)}</strong></p>
  <h2>최근 청소 정보</h2>
  <ul>
    {links}
  </ul>
</main>
</body>
</html>
"""
    write_text(DEPLOYS_DIR / "index.html", content)


# ============================================================
# 사이트맵, RSS, robots
# ============================================================

def xml_urlset(urls: Iterable[str], priority: str) -> str:
    today = datetime.now().strftime("%Y-%m-%d")
    blocks = []
    for url in urls:
        blocks.append(
            "  <url>\n"
            f"    <loc>{esc(url)}</loc>\n"
            f"    <lastmod>{today}</lastmod>\n"
            "    <changefreq>weekly</changefreq>\n"
            f"    <priority>{priority}</priority>\n"
            "  </url>"
        )
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n'
        + "\n".join(blocks)
        + "\n</urlset>\n"
    )


def xml_sitemap_index(filenames: Iterable[str]) -> str:
    today = datetime.now().strftime("%Y-%m-%d")
    blocks = []
    for filename in filenames:
        url = f"{SITE_URL.rstrip('/')}/{filename}"
        blocks.append(
            "  <sitemap>\n"
            f"    <loc>{esc(url)}</loc>\n"
            f"    <lastmod>{today}</lastmod>\n"
            "  </sitemap>"
        )
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<sitemapindex xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n'
        + "\n".join(blocks)
        + "\n</sitemapindex>\n"
    )


def generate_sitemaps(pages: Sequence[Page]) -> list[str]:
    sitemap_files = ["sitemap-main.xml"]
    write_text(
        DEPLOYS_DIR / "sitemap-main.xml",
        xml_urlset([SITE_URL.rstrip("/") + "/"], "1.0"),
    )

    urls = [page.url for page in pages]
    for start in range(0, len(urls), MAX_URLS_PER_SITEMAP):
        number = start // MAX_URLS_PER_SITEMAP + 1
        filename = f"sitemap-pages-{number}.xml"
        sitemap_files.append(filename)
        write_text(
            DEPLOYS_DIR / filename,
            xml_urlset(urls[start:start + MAX_URLS_PER_SITEMAP], "0.8"),
        )

    # 예전 분할 사이트맵이 더 많이 남아 있으면 삭제
    current = set(sitemap_files)
    for old_file in DEPLOYS_DIR.glob("sitemap-pages-*.xml"):
        if old_file.name not in current:
            old_file.unlink()

    write_text(
        DEPLOYS_DIR / "sitemap.xml",
        xml_sitemap_index(sitemap_files),
    )
    return sitemap_files


def generate_rss(pages: Sequence[Page]) -> None:
    now = datetime.now(timezone.utc).strftime("%a, %d %b %Y %H:%M:%S +0000")
    selected = list(pages)[-RSS_MAX_ITEMS:][::-1]

    items = []
    for page in selected:
        items.append(
            "    <item>\n"
            f"      <title>{esc(page.keyword)}</title>\n"
            f"      <link>{esc(page.url)}</link>\n"
            f"      <guid isPermaLink=\"true\">{esc(page.url)}</guid>\n"
            f"      <description>{esc(make_description(page))}</description>\n"
            f"      <pubDate>{now}</pubDate>\n"
            "    </item>"
        )

    rss = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<rss version="2.0">\n'
        "  <channel>\n"
        f"    <title>{esc(BRAND_NAME)}</title>\n"
        f"    <link>{esc(SITE_URL)}</link>\n"
        "    <description>지역별 청소 작업 정보</description>\n"
        "    <language>ko</language>\n"
        f"    <lastBuildDate>{now}</lastBuildDate>\n"
        + "\n".join(items)
        + "\n  </channel>\n"
        "</rss>\n"
    )
    write_text(DEPLOYS_DIR / "rss.xml", rss)


def generate_robots() -> None:
    write_text(
        DEPLOYS_DIR / "robots.txt",
        "User-agent: *\n"
        "Allow: /\n\n"
        f"Sitemap: {SITE_URL.rstrip('/')}/sitemap.xml\n",
    )


def validate_xml(paths: Iterable[Path]) -> None:
    for path in paths:
        ET.parse(path)


# ============================================================
# 검증
# ============================================================

META_DESCRIPTION_RE = re.compile(
    r'<meta\s+name=["\']description["\']\s+content=["\'](.*?)["\']',
    re.I | re.S,
)
TITLE_RE = re.compile(r"<title>(.*?)</title>", re.I | re.S)
H1_RE = re.compile(r"<h1[^>]*>(.*?)</h1>", re.I | re.S)
CANONICAL_RE = re.compile(
    r'<link\s+rel=["\']canonical["\']\s+href=["\'](.*?)["\']',
    re.I | re.S,
)
ROBOTS_RE = re.compile(
    r'<meta\s+name=["\']robots["\']\s+content=["\'](.*?)["\']',
    re.I | re.S,
)
INTERNAL_LINK_RE = re.compile(r'href=["\']/([^"\']+)/?["\']', re.I)


def first_match(pattern: re.Pattern[str], content: str) -> str:
    match = pattern.search(content)
    if not match:
        return ""
    return clean(html.unescape(strip_tags(match.group(1))))


def verify_pages(pages: Sequence[Page], min_links: int = 5) -> tuple[list[VerifyResult], dict]:
    results: list[VerifyResult] = []
    title_map: dict[str, list[int]] = {}
    desc_map: dict[str, list[int]] = {}
    canonical_map: dict[str, list[int]] = {}

    for index, page in enumerate(pages):
        issues: list[str] = []
        score = 100

        if not page.path.exists():
            results.append(
                VerifyResult(
                    keyword=page.keyword,
                    slug=page.slug,
                    status="ERROR",
                    score=0,
                    issues="index.html 없음",
                    title="",
                    description="",
                    canonical="",
                    h1="",
                    internal_links=0,
                )
            )
            continue

        content = page.path.read_text(encoding="utf-8", errors="replace")
        title = first_match(TITLE_RE, content)
        description = first_match(META_DESCRIPTION_RE, content)
        h1 = first_match(H1_RE, content)
        canonical = first_match(CANONICAL_RE, content)
        robots = first_match(ROBOTS_RE, content)
        links = INTERNAL_LINK_RE.findall(content)
        unique_links = set(links)

        def fail(message: str, points: int) -> None:
            nonlocal score
            issues.append(message)
            score -= points

        if not title:
            fail("title 없음", 15)
        elif not 20 <= len(title) <= 70:
            fail("title 길이 확인", 4)

        if not description:
            fail("description 없음", 12)
        elif not 50 <= len(description) <= 170:
            fail("description 길이 확인", 3)

        if not h1:
            fail("H1 없음", 12)
        elif normalize_compare(h1) != normalize_compare(page.keyword):
            fail("H1 키워드 불일치", 4)

        if not canonical:
            fail("canonical 없음", 12)
        elif canonical.rstrip("/") != page.url.rstrip("/"):
            fail("canonical 주소 불일치", 10)

        if "index" not in robots.lower() or "follow" not in robots.lower():
            fail("robots 확인", 8)

        if "application/ld+json" not in content:
            fail("JSON-LD 없음", 10)

        # 실제 템플릿 플레이스홀더만 검사합니다.
        # JSON-LD의 중첩 객체는 정상적으로 "}}"가 포함될 수 있으므로
        # 단순 문자열 포함 여부로 검사하면 모든 페이지가 오탐 처리됩니다.
        template_token_re = re.compile(
            r"\{\{\s*[A-Za-z_가-힣][^{}]{0,120}\s*\}\}",
            re.S,
        )
        if template_token_re.search(content):
            fail("템플릿 토큰 잔존", 10)

        if len(unique_links) < min_links:
            fail(f"내부링크 부족({len(unique_links)})", 6)

        if len(links) != len(unique_links):
            fail("중복 내부링크", 2)

        title_key = normalize_compare(title)
        desc_key = normalize_compare(description)
        canonical_key = canonical.strip().lower()

        if title_key:
            title_map.setdefault(title_key, []).append(index)
        if desc_key:
            desc_map.setdefault(desc_key, []).append(index)
        if canonical_key:
            canonical_map.setdefault(canonical_key, []).append(index)

        status = "PASS" if score >= 90 and not issues else "CHECK"
        results.append(
            VerifyResult(
                keyword=page.keyword,
                slug=page.slug,
                status=status,
                score=max(0, score),
                issues=" | ".join(issues),
                title=title,
                description=description,
                canonical=canonical,
                h1=h1,
                internal_links=len(unique_links),
            )
        )

    # 전체 중복 검사
    for mapping, label, penalty in (
        (title_map, "중복 title", 12),
        (desc_map, "중복 description", 8),
        (canonical_map, "중복 canonical", 20),
    ):
        for indexes in mapping.values():
            if len(indexes) <= 1:
                continue
            for idx in indexes:
                if idx >= len(results):
                    continue
                result = results[idx]
                result.score = max(0, result.score - penalty)
                result.issues = (result.issues + " | " if result.issues else "") + label
                result.status = "ERROR" if label == "중복 canonical" else "CHECK"

    pass_count = sum(1 for item in results if item.status == "PASS")
    check_count = sum(1 for item in results if item.status == "CHECK")
    error_count = sum(1 for item in results if item.status == "ERROR")
    average = round(
        sum(item.score for item in results) / len(results),
        2,
    ) if results else 0.0

    summary = {
        "checked_at": datetime.now().isoformat(timespec="seconds"),
        "total": len(results),
        "pass": pass_count,
        "check": check_count,
        "error": error_count,
        "average_score": average,
    }
    return results, summary


def save_verify_reports(results: Sequence[VerifyResult], summary: dict) -> None:
    with VERIFY_REPORT.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=[
                "keyword", "slug", "status", "score", "issues",
                "title", "description", "canonical", "h1", "internal_links",
            ],
        )
        writer.writeheader()
        for result in results:
            writer.writerow(result.__dict__)

    VERIFY_SUMMARY.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


# ============================================================
# Git 배포
# ============================================================

def run_command(command: Sequence[str], label: str, allow_nothing: bool = False) -> int:
    log("\n" + "=" * 70)
    log(label)
    log("$ " + " ".join(command))
    log("=" * 70)

    try:
        result = subprocess.run(
            list(command),
            cwd=BASE_DIR,
            text=True,
            encoding="utf-8",
            errors="replace",
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
        )
    except FileNotFoundError:
        log(f"{command[0]} 명령을 찾지 못했습니다.")
        return 127

    output = result.stdout or ""
    if output.strip():
        print(output)

    if result.returncode != 0:
        if allow_nothing and "nothing to commit" in output.lower():
            return 0
        return result.returncode
    return 0


def git_deploy(message: str, dry_run: bool = False) -> int:
    commands = [
        ["git", "add", "."],
        ["git", "commit", "-m", message],
        ["git", "push", "origin", "master"],
    ]

    for command in commands:
        if dry_run:
            log("DRY RUN: " + " ".join(command))
            continue

        code = run_command(
            command,
            "Git 배포",
            allow_nothing=(command[1] == "commit"),
        )
        if code:
            return code
    return 0


# ============================================================
# 실행 보고서
# ============================================================

def save_deploy_report(
    total: int,
    created: int,
    skipped: int,
    linked: int,
    sitemap_files: int,
    verify_summary: dict,
) -> None:
    with DEPLOY_REPORT.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=[
                "generated_at", "site_url", "keyword_count",
                "created_pages", "skipped_pages", "linked_pages",
                "sitemap_files", "verify_pass", "verify_check",
                "verify_error", "verify_score",
            ],
        )
        writer.writeheader()
        writer.writerow(
            {
                "generated_at": datetime.now().isoformat(timespec="seconds"),
                "site_url": SITE_URL,
                "keyword_count": total,
                "created_pages": created,
                "skipped_pages": skipped,
                "linked_pages": linked,
                "sitemap_files": sitemap_files,
                "verify_pass": verify_summary.get("pass", 0),
                "verify_check": verify_summary.get("check", 0),
                "verify_error": verify_summary.get("error", 0),
                "verify_score": verify_summary.get("average_score", 0),
            }
        )


# ============================================================
# 실행 옵션
# ============================================================

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="V6.1 키워드 생성 제외 통합 원클릭 엔진"
    )
    parser.add_argument("--limit", type=int, default=0, help="테스트할 페이지 수")
    parser.add_argument("--force", action="store_true", help="기존 페이지도 다시 생성")
    parser.add_argument("--no-git", action="store_true", help="Git 배포하지 않음")
    parser.add_argument("--dry-run", action="store_true", help="Git 명령만 출력")
    parser.add_argument("--message", default="", help="Git commit 메시지")
    parser.add_argument("--max-links", type=int, default=DEFAULT_MAX_LINKS, help="페이지당 내부링크 수")
    parser.add_argument("--min-links", type=int, default=5, help="검증 시 최소 내부링크 수")
    parser.add_argument("--allow-check", action="store_true", help="CHECK 페이지가 있어도 Git 배포")
    parser.add_argument("--allow-error", action="store_true", help="ERROR 페이지가 있어도 Git 배포")
    parser.add_argument("--clean", action="store_true", help="deploys 폴더 삭제 후 새로 생성")
    return parser.parse_args()


# ============================================================
# 메인
# ============================================================

def main() -> int:
    args = parse_args()

    log("=" * 70)
    log("V6.1 통합 원클릭 엔진")
    log("키워드 생성 단계는 포함하지 않습니다.")
    log(f"사용 키워드: {KEYWORDS_FILE}")
    log(f"배포 주소: {SITE_URL}")
    log("=" * 70)

    if args.clean and DEPLOYS_DIR.exists():
        log("deploys 폴더를 삭제합니다.")
        shutil.rmtree(DEPLOYS_DIR)

    DEPLOYS_DIR.mkdir(parents=True, exist_ok=True)

    try:
        pages = read_pages(limit=max(0, args.limit))
    except Exception as exc:
        log(f"키워드 읽기 실패: {exc}")
        return 1

    if not pages:
        log("처리할 키워드가 없습니다.")
        return 1

    log(f"\n키워드 불러오기 완료: {len(pages):,}개")

    log("\n[1/6] 신규 페이지 및 SEO 생성")
    created, skipped = generate_pages(pages, force=args.force)
    log(f"신규 생성: {created:,}개")
    log(f"기존 유지: {skipped:,}개")

    log("\n[2/6] 내부링크 적용")
    linked = apply_internal_links(
        pages,
        max_links=max(1, args.max_links),
    )
    log(f"내부링크 반영: {linked:,}개")

    log("\n[3/6] 메인 페이지 생성")
    generate_home(pages)

    log("\n[4/6] 사이트맵 · RSS · robots 생성")
    sitemap_files = generate_sitemaps(pages)
    generate_rss(pages)
    generate_robots()

    xml_paths = [
        DEPLOYS_DIR / "sitemap.xml",
        DEPLOYS_DIR / "rss.xml",
        *[DEPLOYS_DIR / filename for filename in sitemap_files],
    ]
    try:
        validate_xml(xml_paths)
    except ET.ParseError as exc:
        log(f"XML 문법 오류: {exc}")
        return 1

    log(f"사이트맵 파일: {len(sitemap_files) + 1:,}개")

    log("\n[5/6] 전체 품질검사")
    results, verify_summary = verify_pages(
        pages,
        min_links=max(0, args.min_links),
    )
    save_verify_reports(results, verify_summary)

    log(f"전체: {verify_summary['total']:,}")
    log(f"PASS: {verify_summary['pass']:,}")
    log(f"CHECK: {verify_summary['check']:,}")
    log(f"ERROR: {verify_summary['error']:,}")
    log(f"최종 점수: {verify_summary['average_score']} / 100")

    save_deploy_report(
        total=len(pages),
        created=created,
        skipped=skipped,
        linked=linked,
        sitemap_files=len(sitemap_files) + 1,
        verify_summary=verify_summary,
    )

    if verify_summary["error"] and not args.allow_error:
        log("\nERROR 페이지가 있어 Git 배포를 중단했습니다.")
        log("verify_report_v6_1.csv를 확인하세요.")
        log("강제로 진행하려면 --allow-error 옵션을 사용하세요.")
        return 2

    if verify_summary["check"] and not args.allow_check:
        log("\nCHECK 페이지가 있어 Git 배포를 중단했습니다.")
        log("먼저 테스트 결과를 확인한 뒤 --allow-check로 진행하세요.")
        if not args.no_git:
            return 3

    log("\n[6/6] Git 배포")
    if args.no_git:
        log("Git 배포를 건너뛰었습니다.")
    else:
        message = args.message or (
            f"V6.1 integrated deploy "
            f"{datetime.now().strftime('%Y-%m-%d %H:%M')}"
        )
        code = git_deploy(message, dry_run=args.dry_run)
        if code:
            log(f"Git 배포 실패: 종료 코드 {code}")
            return code

    log("\n" + "=" * 70)
    log("전체 작업 완료")
    log(f"사이트: {SITE_URL}")
    log(f"사이트맵: {SITE_URL.rstrip('/')}/sitemap.xml")
    log(f"RSS: {SITE_URL.rstrip('/')}/rss.xml")
    log(f"검증 리포트: {VERIFY_REPORT.name}")
    log(f"배포 리포트: {DEPLOY_REPORT.name}")
    log("=" * 70)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
